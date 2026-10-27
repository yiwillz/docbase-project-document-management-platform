"""项目结构标准 Excel 模板：生成空白模板、解析、导入（PRD F2）。

这是平台自己定义的固定格式，不解析客户的 MDL。三张表：
- 里程碑：一行一个，行的顺序就是 M1、M2…
- 文件种类：一行一个，同一里程碑内按行的顺序自动编号；"编号"列只用来核对
- 模块分区（可选）：某个文件种类在某个模块下的 MDL 状态和备注

表头中英文都认。解析不抛异常，问题都放进 errors / warnings（带表名、行号、列名）。
只能导入到还没有结构的项目；有任何错误整批不写。
"""

import io
import re
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import exists, select
from sqlalchemy.orm import Session

from app.models import DocumentTypeOrigin, Milestone, Module
from app.services import structure
from app.services.errors import Conflict
from app.services.history import Actor
from app.services.projects import get_writable_project

# ---------------------------------------------------------------- 模板定义


@dataclass(frozen=True)
class Column:
    key: str
    zh: str
    en: str
    kind: str = "text"  # text / date / bool / list / origin
    required: bool = False
    width: int = 18
    note: str = ""


@dataclass(frozen=True)
class SheetDef:
    key: str
    zh: str
    en: str
    columns: tuple[Column, ...]
    optional: bool = False


MILESTONES = SheetDef("milestones", "里程碑", "Milestones", (
    Column("name", "名称", "Name", required=True, width=36, note="行的顺序就是 M1、M2…"),
    Column("goal", "目标", "Goal", width=40),
    Column("description", "说明", "Description", width=40),
    Column("remarks", "备注", "Remarks", width=30),
    Column("due_date", "截止日期", "Due Date", kind="date", width=14),
))

DOCUMENT_TYPES = SheetDef("document_types", "文件种类", "Document Types", (
    Column("milestone", "里程碑", "Milestone", required=True, width=10, note="填 M3 或 3"),
    Column("submittal_no", "编号", "No.", width=10,
           note="可不填；填了必须等于按行顺序算出的编号，用来核对"),
    Column("name", "名称", "Name", required=True, width=40),
    Column("category", "分类", "Category", width=22),
    Column("per_equipment", "按设备", "Per Equipment", kind="bool", width=10, note="是 / 否"),
    Column("modules", "模块", "Modules", kind="list", width=30, note="多个用逗号分隔"),
    Column("due_date", "截止日期", "Due Date", kind="date", width=14),
    Column("content", "应含内容", "Content", width=40),
    Column("description", "说明", "Description", width=40),
    Column("purpose", "目的", "Purpose", width=40),
    Column("mdl_status_text", "MDL状态", "MDL Status", width=24),
    Column("mdl_remarks", "MDL备注", "MDL Remarks", width=30),
    Column("mdl_submitted_date", "提交日期", "Submitted Date", kind="date", width=14),
    Column("origin", "要求类别", "Requirement", kind="origin", width=12,
           note="合同规定 / 客户要求 / 内部要求；不填为合同规定"),
    Column("required_date", "要求日期", "Required Date", kind="date", width=14,
           note="合同签订日 / 客户提出日 / 内部决定日；不知道就空着"),
    Column("basis", "依据", "Basis", width=22,
           note="合同规定填合同条款；客户要求填客户提出要求的函件、邮件、会议纪要编号"),
    Column("changed_date", "变更日期", "Changed Date", kind="date", width=14,
           note="要求内容后来变更的日期"),
    Column("change_ref", "变更依据", "Change Ref", width=22, note="如客户函件号、邮件"),
    Column("in_client_mdl", "列入客户MDL", "In Client MDL", kind="bool", width=12,
           note="是 / 否；不填为是；内部要求一律为否"),
))

PARTITIONS = SheetDef("partitions", "模块分区", "Module Partitions", (
    Column("submittal_no", "编号", "No.", required=True, width=10, note="对应文件种类的编号"),
    Column("module", "模块", "Module", required=True, width=22),
    Column("mdl_status_text", "MDL状态", "MDL Status", width=24),
    Column("mdl_remarks", "MDL备注", "MDL Remarks", width=30),
), optional=True)

SHEETS = (MILESTONES, DOCUMENT_TYPES, PARTITIONS)

TRUE_WORDS = {"是", "y", "yes", "true", "1", "√", "✓", "✔", "x"}
FALSE_WORDS = {"否", "n", "no", "false", "0"}
ERROR_VALUES = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}
_LIST_SPLIT = re.compile(r"[,，;；\n]+")
# 要求类别的各种写法 → 代码
ORIGIN_WORDS = {
    DocumentTypeOrigin.CONTRACT: {"合同规定", "合同", "contract"},
    DocumentTypeOrigin.CLIENT: {"客户要求", "客户", "追加", "客户追加", "client"},
    DocumentTypeOrigin.INTERNAL: {"内部要求", "内部", "internal"},
}


# ---------------------------------------------------------------- 生成模板


def build_workbook(lang: str = "zh") -> Workbook:
    """空白模板：三张数据表 + 一张说明表。"""
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    required_fill = PatternFill("solid", fgColor="C00000")
    for sheet in SHEETS:
        ws = wb.create_sheet(sheet.zh if lang == "zh" else sheet.en)
        for col, column in enumerate(sheet.columns, start=1):
            cell = ws.cell(1, col, column.zh if lang == "zh" else column.en)
            cell.font = header_font
            cell.fill = required_fill if column.required else header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = column.width
            if column.kind in ("bool", "origin"):
                choices = {("bool", "zh"): '"是,否"', ("bool", "en"): '"Yes,No"',
                           ("origin", "zh"): '"合同规定,客户要求,内部要求"',
                           ("origin", "en"): '"contract,client,internal"'}
                dv = DataValidation(type="list", formula1=choices[(column.kind, lang)],
                                    allow_blank=True)
                dv.add(f"{cell.column_letter}2:{cell.column_letter}5000")
                ws.add_data_validation(dv)
        ws.freeze_panes = "A2"

    guide = wb.create_sheet("说明" if lang == "zh" else "Guide")
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 22
    guide.column_dimensions["C"].width = 70
    lines = [
        ("项目结构导入模板", "", ""),
        ("", "", ""),
        ("1. 只能导入到还没有结构的项目；有任何错误整批不导入。", "", ""),
        ("2. 编号全部由平台自动生成：里程碑按行顺序为 M1、M2…；文件种类在同一里程碑内按行顺序为 x.1、x.2…", "", ""),
        ("3. 红色表头为必填列。", "", ""),
        ("", "", ""),
        ("表", "列", "说明"),
    ]
    for sheet in SHEETS:
        for column in sheet.columns:
            lines.append((sheet.zh, column.zh, column.note))
    for row in lines:
        guide.append(row)
    guide["A1"].font = Font(bold=True, size=14)
    for cell in guide[7]:
        cell.font = Font(bold=True)
    return wb


def build_template(lang: str = "zh") -> bytes:
    buffer = io.BytesIO()
    build_workbook(lang).save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- 解析


@dataclass
class Issue:
    code: str
    message: str
    sheet: str | None = None
    row: int | None = None
    column: str | None = None


@dataclass
class ParsedStructure:
    milestones: list[dict[str, Any]] = field(default_factory=list)
    document_types: list[dict[str, Any]] = field(default_factory=list)
    partitions: list[dict[str, Any]] = field(default_factory=list)
    modules: list[str] = field(default_factory=list)
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def summary(self) -> dict[str, int]:
        return {
            "milestones": len(self.milestones),
            "document_types": len(self.document_types),
            "per_equipment": sum(1 for d in self.document_types if d.get("per_equipment")),
            "per_module": sum(1 for d in self.document_types if d["modules"]),
            "modules": len(self.modules),
            "partitions": sum(len(d["modules"]) for d in self.document_types),
        }

    def to_json(self) -> dict[str, Any]:
        """存进任务表用：日期转成文字。"""
        def plain(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
            return [{k: v.isoformat() if isinstance(v, date) else v for k, v in r.items()}
                    for r in rows]
        return {
            "milestones": plain(self.milestones),
            "document_types": plain(self.document_types),
            "partitions": plain(self.partitions),
            "modules": self.modules,
            "errors": [asdict(i) for i in self.errors],
            "warnings": [asdict(i) for i in self.warnings],
        }


def _norm_header(text: Any) -> str:
    return re.sub(r"\s+", "", str(text or "")).casefold()


def _find_sheet(wb, sheet: SheetDef):
    wanted = {_norm_header(sheet.zh), _norm_header(sheet.en)}
    for ws in wb.worksheets:
        if _norm_header(ws.title) in wanted:
            return ws
    return None


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float):
        value = repr(value) if not value.is_integer() else str(int(value))
    elif isinstance(value, (datetime, date)):
        value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = " ".join(str(value).replace("\xa0", " ").split())
    return None if not text or text in ERROR_VALUES else text


def _convert(column: Column, value: Any) -> Any:
    """按列类型转换；格式不对抛 ValueError（由调用方记成错误）。"""
    if column.kind == "date":
        if value is None or isinstance(value, str) and not value.strip():
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = _text(value)
        if text is None:
            return None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
        raise ValueError(f"日期格式不对：{text}（应为 2026-03-01）")
    text = _text(value)
    if column.kind == "bool":
        if text is None:
            return None
        if text.casefold() in TRUE_WORDS:
            return True
        if text.casefold() in FALSE_WORDS:
            return False
        raise ValueError(f"应填 是 / 否：{text}")
    if column.kind == "origin":
        if text is None:
            return None
        word = text.casefold()
        for origin, words in ORIGIN_WORDS.items():
            if word in words:
                return origin.value
        raise ValueError(f"要求类别应填 合同规定 / 客户要求 / 内部要求：{text}")
    if column.kind == "list":
        return [] if text is None else [p.strip() for p in _LIST_SPLIT.split(text) if p.strip()]
    return text


def _read_sheet(ws, sheet: SheetDef, result: ParsedStructure) -> list[dict[str, Any]]:
    """读一张表：第 1 行是表头，从第 2 行起是数据；整行为空的跳过。"""
    headers = {}
    for col in range(1, ws.max_column + 1):
        key = _norm_header(ws.cell(1, col).value)
        if not key:
            continue
        match = next((c for c in sheet.columns if key in (_norm_header(c.zh), _norm_header(c.en))),
                     None)
        if match is None:
            result.warnings.append(Issue("template.unknown_column",
                                         f"不认识的列 {ws.cell(1, col).value}，已忽略",
                                         sheet.zh, 1, str(ws.cell(1, col).value)))
        else:
            headers[col] = match
    for column in sheet.columns:
        if column.required and column not in headers.values():
            result.errors.append(Issue("template.missing_column", f"缺少必填列：{column.zh}",
                                       sheet.zh, 1, column.zh))
    if any(i.code == "template.missing_column" and i.sheet == sheet.zh for i in result.errors):
        return []

    rows = []
    for r in range(2, ws.max_row + 1):
        raw = {col: ws.cell(r, col).value for col in headers}
        if all(_text(v) is None for v in raw.values()):
            continue
        row: dict[str, Any] = {"row": r}
        for col, column in headers.items():
            try:
                row[column.key] = _convert(column, raw[col])
            except ValueError as exc:
                result.errors.append(Issue("template.bad_value", str(exc), sheet.zh, r, column.zh))
                row[column.key] = None
        for column in sheet.columns:
            if column.required and not row.get(column.key):
                result.errors.append(Issue("template.required", f"{column.zh} 不能为空",
                                           sheet.zh, r, column.zh))
        rows.append(row)
    return rows


def parse_template(data: bytes) -> ParsedStructure:
    result = ParsedStructure()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # 不是 Excel、文件损坏等
        result.errors.append(Issue("template.unreadable", f"无法读取 Excel：{exc}"))
        return result

    sheets = {}
    for sheet in SHEETS:
        ws = _find_sheet(wb, sheet)
        if ws is None and not sheet.optional:
            result.errors.append(Issue("template.missing_sheet",
                                       f"缺少工作表：{sheet.zh}（或 {sheet.en}）", sheet.zh))
        sheets[sheet.key] = _read_sheet(ws, sheet, result) if ws is not None else []
    if not result.ok:
        return result

    # 里程碑：行顺序即编号
    for number, row in enumerate(sheets["milestones"], start=1):
        result.milestones.append({**row, "number": number})

    # 文件种类：按所属里程碑、行顺序自动编号；"编号"列只用来核对
    seqs: dict[int, int] = {}
    by_number: dict[str, dict[str, Any]] = {}
    for row in sheets["document_types"]:
        ref = row.get("milestone") or ""
        match = re.fullmatch(r"\D*(\d+)", ref)
        ms = int(match.group(1)) if match else None
        if ms is None or not 1 <= ms <= len(result.milestones):
            result.errors.append(Issue("template.unknown_milestone",
                                       f"里程碑 {ref} 不存在（里程碑表共 {len(result.milestones)} 行）",
                                       DOCUMENT_TYPES.zh, row["row"], DOCUMENT_TYPES.columns[0].zh))
            continue
        seqs[ms] = seqs.get(ms, 0) + 1
        number = f"{ms}.{seqs[ms]}"
        if row.get("submittal_no") and row["submittal_no"] != number:
            result.errors.append(Issue(
                "template.number_mismatch",
                f"编号 {row['submittal_no']} 与按行顺序算出的 {number} 不一致",
                DOCUMENT_TYPES.zh, row["row"], DOCUMENT_TYPES.columns[1].zh,
            ))
        doc = {**row, "milestone": ms, "submittal_no": number,
               "modules": list(dict.fromkeys(row.get("modules") or [])), "partitions": {}}
        result.document_types.append(doc)
        by_number[number] = doc

    # 模块分区：按编号挂到文件种类；模块不在"模块"列里的自动补上
    for row in sheets["partitions"]:
        doc = by_number.get(row.get("submittal_no") or "")
        if doc is None:
            result.errors.append(Issue("template.unknown_document_type",
                                       f"编号 {row.get('submittal_no')} 在文件种类表里不存在",
                                       PARTITIONS.zh, row["row"], PARTITIONS.columns[0].zh))
            continue
        module = row["module"]
        if module in doc["partitions"]:
            result.errors.append(Issue("template.duplicate_partition",
                                       f"{doc['submittal_no']} 的模块 {module} 重复",
                                       PARTITIONS.zh, row["row"], PARTITIONS.columns[1].zh))
            continue
        if module not in doc["modules"]:
            doc["modules"].append(module)
        doc["partitions"][module] = {"mdl_status_text": row.get("mdl_status_text"),
                                     "mdl_remarks": row.get("mdl_remarks")}
        result.partitions.append({**row})

    result.modules = list(dict.fromkeys(m for d in result.document_types for m in d["modules"]))
    if not result.milestones:
        result.warnings.append(Issue("template.empty", "里程碑表是空的", MILESTONES.zh))
    return result


# ---------------------------------------------------------------- 导入


def _as_date(value: Any) -> date | None:
    return date.fromisoformat(value) if isinstance(value, str) else value


def import_structure(session: Session, project_id: int, parsed: dict[str, Any], actor: Actor,
                     *, source: str | None = None) -> dict[str, int]:
    """把解析结果（ParsedStructure.to_json() 的内容）写进项目。目标项目必须还没有结构。"""
    project = get_writable_project(session, project_id)
    if parsed.get("errors"):
        raise Conflict("structure_import.has_errors", "模板里还有错误，不能导入",
                       errors=len(parsed["errors"]))
    if session.scalar(select(exists().where(Milestone.project_id == project.id))) or \
            session.scalar(select(exists().where(Module.project_id == project.id))):
        raise Conflict("structure.target_not_empty",
                       f"项目 {project.code} 已有结构，模板只能导入到空项目", project_id=project.id)

    reason = f"从标准模板导入（{source}）" if source else "从标准模板导入"
    milestone_ids: dict[int, int] = {}
    for row in parsed["milestones"]:
        milestone = structure.create_milestone(session, project.id, {
            "name": row.get("name"), "goal": row.get("goal"),
            "description": row.get("description"), "remarks": row.get("remarks"),
            "due_date": _as_date(row.get("due_date")),
        }, actor)
        milestone_ids[row["number"]] = milestone.id

    modules = {name: structure.create_module(session, project.id, {"name": name}, actor).id
               for name in parsed["modules"]}

    doc_fields = ("name", "category", "per_equipment", "content", "description", "purpose",
                  "mdl_status_text", "mdl_remarks", "origin", "basis", "change_ref",
                  "in_client_mdl")
    counts = {"milestones": len(milestone_ids), "document_types": 0, "modules": len(modules),
              "partitions": 0}
    for number, milestone_id in milestone_ids.items():
        docs = [d for d in parsed["document_types"] if d["milestone"] == number]
        if not docs:
            continue
        rows = [{**{f: d.get(f) for f in doc_fields},
                 "due_date": _as_date(d.get("due_date")),
                 "mdl_submitted_date": _as_date(d.get("mdl_submitted_date")),
                 "required_date": _as_date(d.get("required_date")),
                 "changed_date": _as_date(d.get("changed_date"))} for d in docs]
        created = structure.bulk_create_document_types(session, project.id, milestone_id, rows,
                                                       actor)
        for doc, row in zip(created, docs, strict=True):
            # 目标项目是空的，自动编号必然和解析时算出的一致；不一致说明程序有问题
            if doc.submittal_no != row["submittal_no"]:
                raise RuntimeError(f"编号不一致：{doc.submittal_no} ≠ {row['submittal_no']}")
            if row["modules"]:
                structure.set_document_type_modules(
                    session, project.id, doc.id, actor,
                    module_ids=[modules[m] for m in row["modules"]], reason=reason)
                for module, values in row["partitions"].items():
                    if any(values.values()):
                        structure.update_partition(session, project.id, doc.id, modules[module],
                                                   values, actor, reason=reason)
                counts["partitions"] += len(row["modules"])
        counts["document_types"] += len(created)
    return counts
