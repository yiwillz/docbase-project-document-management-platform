"""变更记录导出 Excel（PRD F10.2：按项目、对象、人、时间查询和导出）。

每个改动的字段一行：时间、操作人、对象类型、对象、动作、字段、修改前、修改后、原因。
筛选条件和变更记录页一样；最多 MAX_ROWS 行，超了在表头注明只导出了最新的部分。
"""

import io
import json
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.services import mdl_export as mx
from app.services.change_labels import labels, summary_kind
from app.services.history import query_change_logs
from app.services.projects import get_project

MAX_ROWS = 50_000
PAGE = 2_000

ENTITY = {
    "zh": {"projects": "项目", "milestones": "里程碑", "modules": "模块", "document_types": "文件种类",
           "document_type_modules": "模块分区", "document_type_equipment": "种类对应设备", "equipment": "设备",
           "equipment_types": "设备种类", "equipment_units": "单台", "deliverables": "应交项",
           "transmittals": "传送单", "transmittal_items": "传送单明细", "review_codes": "审核代码",
           "status_mappings": "状态映射", "mdl_versions": "MDL 版本", "files": "文件", "versions": "文件版本",
           "baselines": "阶段基线", "project_members": "项目成员", "users": "用户", "import_batches": "存量导入", "module_equipment_types": "模块对应设备种类"},
    "en": {"projects": "Project", "milestones": "Milestone", "modules": "Module", "document_types": "Document type",
           "document_type_modules": "Module partition", "document_type_equipment": "Type–equipment link",
           "equipment": "Equipment", "equipment_types": "Equipment type", "equipment_units": "Unit",
           "deliverables": "Deliverable", "transmittals": "Transmittal", "transmittal_items": "Transmittal item",
           "review_codes": "Review code", "status_mappings": "Status mapping", "mdl_versions": "MDL revision",
           "files": "File", "versions": "File version", "baselines": "Baseline", "project_members": "Member",
           "users": "User", "import_batches": "Bulk import", "module_equipment_types": "Module–equipment type"},
}
TEXT = {
    "zh": {"title": "变更记录", "project": "项目", "filters": "筛选", "exported": "导出时间", "all": "全部",
           "truncated": "超过 {n} 行，只导出了最新的 {n} 行；请缩小时间范围分次导出",
           "cols": ("时间", "操作人", "对象类型", "对象", "动作", "字段", "修改前", "修改后", "原因"),
           "actions": {"create": "新建", "update": "修改", "delete": "删除", "restore": "恢复"},
           "summary": {"recalculate": "重算应交项", "purge": "清理回收区", "transmittal_reply": "整张传送单录入回复"}},
    "en": {"title": "CHANGE LOG", "project": "Project", "filters": "Filters", "exported": "Exported", "all": "All",
           "truncated": "More than {n} rows: only the latest {n} were exported; narrow the date range",
           "cols": ("Time", "User", "Object type", "Object", "Action", "Field", "Before", "After", "Reason"),
           "actions": {"create": "Create", "update": "Update", "delete": "Delete", "restore": "Restore"},
           "summary": {"recalculate": "Recalculate deliverables", "purge": "Recycle bin cleanup",
                       "transmittal_reply": "Reply for whole transmittal"}},
}
WIDTHS = (18, 14, 14, 30, 8, 22, 36, 36, 30)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, default=str)


def _object(log: Any, label: str | None, t: dict) -> str:
    """对象一栏：汇总式记录写成"重算应交项"这样的说明（传送单的带上单号）。"""
    kind = summary_kind(log)
    if kind:
        return t["summary"][kind] + (f" {label}" if label else "")
    return label or log.entity_id


def export(session: Session, project_id: int, *, entity_type: str | None = None, actor: str | None = None,
           since: datetime | None = None, until: datetime | None = None, lang: str = "zh") -> bytes:
    project = get_project(session, project_id)
    t, names = TEXT[lang], ENTITY[lang]
    wb = Workbook()
    ws = wb.active
    ws.title = t["title"][:31]
    ws["A1"] = t["title"]
    ws["A1"].font = Font(size=16, bold=True, color=mx.NAVY)
    ws["A2"], ws["B2"] = t["project"], f"{project.name} ({project.code})"
    span = " ~ ".join(d.strftime("%Y-%m-%d %H:%M") if d else "…" for d in (since, until)) if since or until else ""
    filters = " · ".join(filter(None, (names.get(entity_type, entity_type) if entity_type else None, actor, span)))
    ws["A3"], ws["B3"] = t["filters"], filters or t["all"]
    ws["E2"], ws["F2"] = t["exported"], datetime.now().strftime("%Y-%m-%d %H:%M")
    for cell in ("A2", "A3", "E2"):
        ws[cell].font = Font(bold=True, color="595959")

    head = 5
    for col, (title, width) in enumerate(zip(t["cols"], WIDTHS, strict=True), start=1):
        c = ws.cell(head, col, title)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=mx.NAVY), mx.BOX
        ws.column_dimensions[c.column_letter].width = width
    row, offset, truncated = head + 1, 0, False
    while not truncated:
        total, logs = query_change_logs(session, project_id=project_id, entity_type=entity_type, actor=actor,
                                        since=since, until=until, limit=PAGE, offset=offset)
        if not logs:
            break
        offset += len(logs)
        label = labels(session, logs)
        for log in logs:
            fields = sorted((log.changes or {}).items()) or [("", [None, None])]
            if row - head - 1 + len(fields) > MAX_ROWS:
                truncated = True  # 一条记录的字段不拆开：放不下整条就停
                break
            for field, pair in fields:
                before, after = (pair + [None, None])[:2] if isinstance(pair, list) else (None, pair)
                values = (log.created_at.astimezone().replace(tzinfo=None), log.actor_name,
                          names.get(log.entity_type, log.entity_type), _object(log, label.get(log.id), t),
                          t["actions"].get(log.action.value, log.action.value), field, _text(before),
                          _text(after), log.reason or "")
                for col, v in enumerate(values, start=1):
                    c = ws.cell(row, col, v)
                    c.border = mx.BOX
                    c.alignment = Alignment(vertical="top", wrap_text=col in (4, 7, 8, 9))
                ws.cell(row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
                row += 1
        if offset >= total:
            break
    if truncated:
        ws["A4"] = t["truncated"].format(n=MAX_ROWS)
        ws["A4"].font = Font(bold=True, color="C00000")
    ws.auto_filter.ref = f"A{head}:I{max(row - 1, head)}"
    ws.freeze_panes = f"A{head + 1}"
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{head}:{head}"
    wb.properties.creator = f"{mx.PLATFORM} (system generated)"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
