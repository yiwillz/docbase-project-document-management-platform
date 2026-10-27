"""存量导入与清理（PRD F9）：指定目录 → 扫描（算哈希，不复制）→ 分析（去重、识别、版本归并、疑似新设备）→
预演报告 Excel →（改报告导回）→ 执行（复制进存储，规则和平时上传一样）→ 可整批撤销。原目录只读不改。

扫描、执行、撤销都在后台 worker 里做：每处理一个文件记一次断点，暂停、崩溃后接着做。

决定（用户 2026-10-25 确认）：
- 旧版本全部按顺序作为历史版本导入，照常只留最新 3 版（更早的直接进回收区）；
- 只能导入配置里允许的目录（DOCBASE_IMPORT_ROOTS）；
- 不单独做"索引 Excel"：预演报告本身就是这样一张表，改好位置导回即可；
- 整批撤销：已提交给客户的、导入后又被人升过版的不撤，其余删除。
"""

import io
import os
import re
from collections import defaultdict
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import get_settings, import_roots
from app.models import (
    ChangeAction,
    Deliverable,
    DocumentType,
    Equipment,
    File,
    FileDeliverable,
    FileSource,
    ImportBatch,
    ImportItem,
    Job,
    JobStatus,
    Module,
    Version,
)
from app.services import files as files_svc
from app.services import naming
from app.services import recognizer as rec
from app.services.errors import Conflict, Invalid, NotFound
from app.services.history import Actor, emit, record, record_summary, snapshot
from app.services.projects import get_project, get_writable_project
from app.services.storage import get_storage, hash_file
from app.worker import Context, handler

SCAN, EXECUTE, UNDO = "import_scan", "import_execute", "import_undo"
REPORT_KIND = "import_report"
SKIP_FILES = {"thumbs.db", "desktop.ini", ".ds_store"}
SKIP_DIRS = {"$recycle.bin", "system volume information", "__macosx"}
COMMIT_EVERY = 20


# ---------------------------------------------------------------- 批次


def check_source(source_dir: str) -> Path:
    roots = import_roots()
    if not roots:
        raise Invalid("import.no_roots", "没有配置允许导入的目录（DOCBASE_IMPORT_ROOTS），请管理员在 .env 里设置")
    path = Path(source_dir.strip().strip('"')).expanduser()
    try:
        path = path.resolve()
    except OSError as exc:
        raise Invalid("import.bad_dir", f"目录不对：{source_dir}") from exc
    if not any(path == r or path.is_relative_to(r) for r in roots):
        raise Invalid("import.not_allowed", f"只能导入这些目录下的文件：{'；'.join(map(str, roots))}",
                      roots=[str(r) for r in roots])
    if not path.is_dir():
        raise Invalid("import.not_found", f"目录不存在或打不开：{path}", path=str(path))
    return path


def get_batch(session: Session, project_id: int, batch_id: int) -> ImportBatch:
    batch = session.get(ImportBatch, batch_id)
    if batch is None or batch.project_id != project_id:
        raise NotFound("import.not_found", f"导入批次 {batch_id} 不存在", id=batch_id)
    return batch


def _queue(session: Session, batch: ImportBatch, kind: str, actor: Actor) -> Job:
    job = Job(kind=kind, project_id=batch.project_id, status=JobStatus.QUEUED,
              payload={"batch_id": batch.id}, created_by=actor.id)
    session.add(job)
    session.flush()
    batch.job_id = job.id
    return job


def create_batch(session: Session, project_id: int, source_dir: str, actor: Actor, *,
                 description: str | None = None) -> ImportBatch:
    """新建一次导入：排一个扫描任务（后台做）。"""
    get_writable_project(session, project_id)
    path = check_source(source_dir)
    batch = ImportBatch(project_id=project_id, source_dir=str(path),
                        description=(description or "").strip() or None, status="scanning",
                        created_by_name=actor.name)
    session.add(batch)
    session.flush()
    _queue(session, batch, SCAN, actor)
    record(session, batch, ChangeAction.CREATE, None, snapshot(batch), actor)
    return batch


def _job_out(session: Session, job_id: int | None) -> dict[str, Any] | None:
    job = session.get(Job, job_id) if job_id else None
    if job is None:
        return None
    return {"id": job.id, "kind": job.kind, "status": job.status.value, "progress": job.progress,
            "total": job.total, "error": (job.error or "").split("\n")[0] or None}


def batch_out(session: Session, batch: ImportBatch) -> dict[str, Any]:
    return {"id": batch.id, "source_dir": batch.source_dir, "description": batch.description,
            "status": batch.status, "counts": batch.counts or {}, "error": batch.error,
            "created_by": batch.created_by_name, "created_at": batch.created_at,
            "executed_at": batch.executed_at, "undone_at": batch.undone_at,
            "job": _job_out(session, batch.job_id)}


def list_batches(session: Session, project_id: int) -> list[dict[str, Any]]:
    get_project(session, project_id)
    return [batch_out(session, b) for b in session.scalars(
        select(ImportBatch).where(ImportBatch.project_id == project_id).order_by(ImportBatch.id.desc()))]


def set_paused(session: Session, project_id: int, batch_id: int, paused: bool) -> ImportBatch:
    """暂停 / 继续当前的后台任务（扫描、执行、撤销）。"""
    get_writable_project(session, project_id)
    batch = get_batch(session, project_id, batch_id)
    job = session.get(Job, batch.job_id) if batch.job_id else None
    if job is None:
        raise Conflict("import.no_job", "这个批次现在没有后台任务")
    if paused and job.status in (JobStatus.QUEUED, JobStatus.RUNNING):
        job.status = JobStatus.PAUSED
    elif not paused and job.status in (JobStatus.PAUSED, JobStatus.FAILED):
        job.status, job.lease_owner, job.lease_until = JobStatus.QUEUED, None, None
    else:
        raise Conflict("import.bad_state", f"任务现在是 {job.status.value}，不能{'暂停' if paused else '继续'}")
    return batch


# ---------------------------------------------------------------- 扫描


def _walk(root: Path) -> list[Path]:
    out = []
    for dirpath, dirnames, filenames in os.walk(root, followlinks=False):
        dirnames[:] = sorted(d for d in dirnames
                             if not d.startswith(".") and d.casefold() not in SKIP_DIRS)
        for name in sorted(filenames):
            if name.startswith(("~$", ".")) or name.casefold() in SKIP_FILES:
                continue
            out.append(Path(dirpath) / name)
    return out


def scan(session: Session, batch: ImportBatch, ctx: Context | None = None) -> int:
    """遍历目录，每个文件记大小、修改日、哈希（不复制）；已扫过的跳过（断点续跑）。"""
    root = Path(batch.source_dir)
    if not root.is_dir():
        raise Invalid("import.not_found", f"目录不存在或打不开：{root}")
    seen = set(session.scalars(select(ImportItem.rel_path).where(ImportItem.batch_id == batch.id)))
    paths = _walk(root)
    for n, path in enumerate(paths, 1):
        rel = path.relative_to(root).as_posix()
        if rel not in seen:
            try:
                st = path.stat()
                item = ImportItem(batch_id=batch.id, rel_path=rel, size=st.st_size,
                                  modified=datetime.fromtimestamp(st.st_mtime, UTC).date(),
                                  sha256=hash_file(path))
            except OSError as exc:
                item = ImportItem(batch_id=batch.id, rel_path=rel, size=0, sha256="",
                                  state="unreadable", action="skip", error=f"读不了：{exc}")
            session.add(item)
        if n % COMMIT_EVERY == 0:
            session.commit()
            if ctx:
                ctx.heartbeat(progress=n, total=len(paths))
    session.commit()
    if ctx:
        ctx.heartbeat(progress=len(paths), total=len(paths))
    return len(paths)


# ---------------------------------------------------------------- 分析


_TAG_LIKE = re.compile(r"(?<![\w.])([A-Za-z]{0,3}-?\d{2,4}[A-Za-z]?)(?![\w.])")


def _suspected(name: str, folder: str, known_tags: set[str]) -> str | None:
    """像设备编号、但项目里没有这台设备的：如文件夹 "Tool 205"。"""
    for text in (folder, naming.split_ext(name)[0]):
        cleaned = rec.strip_noise(text)
        for m in _TAG_LIKE.finditer(cleaned):
            token = m.group(1)
            if token.casefold() not in known_tags and not re.fullmatch(r"\d{1,2}", token):
                return token
    return None


def _version_key(item: ImportItem) -> tuple:
    name = item.rel_path.rsplit("/", 1)[-1]
    v = naming.parse_version(name)
    value = -1.0
    if v:
        m = re.search(r"(\d+(?:\.\d+)?)\s*$", v) or re.search(r"([A-Za-z])\s*$", v)
        if m:
            text = m.group(1)
            value = float(text) if text[0].isdigit() else float(ord(text.upper()) - 64)
    return (value, naming.parse_date(name) or item.modified or date.min, item.modified or date.min, item.id)


def _name(item: ImportItem) -> str:
    return item.rel_path.rsplit("/", 1)[-1]


def _folder(item: ImportItem) -> str:
    return item.rel_path.rsplit("/", 1)[0] if "/" in item.rel_path else ""


def analyze(session: Session, batch: ImportBatch, *, recognize: bool = True) -> dict[str, int]:
    """去重、识别（只在第一次）、版本归并、算新文件名，更新数量。"""
    project_id = batch.project_id
    project = get_project(session, project_id)
    items = session.scalars(select(ImportItem).where(ImportItem.batch_id == batch.id)
                            .order_by(ImportItem.rel_path)).all()
    if recognize:
        engine = rec.RuleRecognizer(rec.ProjectIndex(session, project_id))
        known_tags = {e.tag.casefold() for e in session.scalars(
            select(Equipment).where(Equipment.project_id == project_id)) if e.tag}
        for item in items:
            if item.state == "unreadable":
                continue
            found = engine.recognize(rec.RecognizeInput(filename=_name(item), folder=_folder(item),
                                                        description=batch.description))
            best = found.best
            item.rec_deliverable_ids = best.deliverable_ids if best else []
            item.rec_level = best.level if best else "low"
            item.rec_reasons = [r.__dict__ for r in best.reasons] if best else []
            item.deliverable_ids = best.deliverable_ids if best and best.level != "low" else []
            item.source = "client" if found.source == "client" else None
            has_equipment = best and any(r.kind in ("tag", "equipment_name", "platform_name")
                                         for r in best.reasons)
            item.suspected_equipment = None if has_equipment else _suspected(
                _name(item), _folder(item), known_tags)
        # 去重：平台已有的不导；本批内容相同的只导一份——留最有用的那份
        # （识别得好的 → 文件名里有版本号 → 有日期 → 名字长的），不管目录里谁排在前面
        by_sha: dict[str, list[ImportItem]] = defaultdict(list)
        for item in items:
            if item.state != "unreadable":
                by_sha[item.sha256].append(item)
        level_rank = {"high": 2, "medium": 1}
        for sha, same in by_sha.items():
            if files_svc.find_duplicate(session, project_id, sha):
                for item in same:
                    item.state, item.action = "existing", "skip"
                continue
            keep = max(same, key=lambda i: (level_rank.get(i.rec_level or "", 0),
                                            bool(naming.parse_version(_name(i))),
                                            bool(naming.parse_date(_name(i))), len(_name(i)), -i.id))
            for item in same:
                if item is keep:
                    item.state, item.duplicate_of = "new", None
                else:
                    item.state, item.duplicate_of, item.action = "duplicate", keep.id, "skip"
        session.flush()

    # 版本归并：同一位置、同一标题（按设备的同一台就算）的一组，按版本号 → 日期 → 修改日排序
    groups: dict[str, list[ImportItem]] = defaultdict(list)
    docs: dict[int, DocumentType] = {}
    for item in items:
        item.group_key = item.version_rank = item.planned_name = None
        if item.state in ("old_version",):
            item.state = "new"
        if item.action != "import" or item.state not in ("new",):
            continue
        ids = sorted(item.deliverable_ids or [])
        title = item.title or _title(session, item, docs)
        # 同一份文件 = 同一位置 + 同一标题 + 同一格式（用户 2026-10-25 确认：一台设备下可以有好几份文件）
        key = ("ids:" + ",".join(map(str, ids)) if ids else "u:" + _folder(item)) + "|" + \
            (item.source or "") + "|" + naming.title_key(title) + "|" + files_svc.ext_of(_name(item))
        groups[key].append(item)
    per_set: dict[str, int] = defaultdict(int)  # 这一批里同一位置有几份（非客户）文件
    for key, members in groups.items():
        if key.startswith("ids:") and members[0].source != "client":
            per_set[key.split("|", 1)[0]] += 1
    for key, members in groups.items():
        members.sort(key=_version_key)
        for rank, item in enumerate(members, 1):
            item.group_key, item.version_rank = key[:500], rank
            if rank < len(members):
                item.state = "old_version"
        _plan_names(session, project, members, batch_siblings=per_set.get(key.split("|", 1)[0], 1) - 1)
    batch.counts = counts(session, batch)
    session.flush()
    return batch.counts


def _title(session: Session, item: ImportItem, docs: dict[int, DocumentType]) -> str:
    """按识别到的位置去掉文件名开头的编号、种类名、设备名后的标题。"""
    ids = item.deliverable_ids or []
    if not ids:
        return naming.extract_title(_name(item))
    target = files_svc.load_target_any(session, 0, ids)
    docs.setdefault(target.doc.id, target.doc)
    return naming.extract_title(_name(item), no=target.doc.submittal_no, type_name=target.doc.name,
                                equipment_names=[e.name for e in target.equipment if e])


def _plan_names(session: Session, project, members: list[ImportItem], *, batch_siblings: int = 0) -> None:
    """一组里每个文件的新文件名：第一份按平台现有的版本往后排，后面的依次加 1。
    按设备的一项会有好几份文件时（平台里已有别的、或这一批里还有别的），名字里加标题。"""
    first = members[0]
    try:
        plan = files_svc.plan(session, project.id, first.deliverable_ids, original_name=_name(first),
                              sha256=first.sha256, modified=first.modified,
                              overrides={k: v for k, v in (("source", first.source),
                                                           ("title", first.title)) if v})
    except (Invalid, NotFound, Conflict) as exc:
        for item in members:
            item.planned_name, item.error = None, exc.message
        return
    target = files_svc.load_target(session, project.id, first.deliverable_ids) \
        if first.deliverable_ids else None
    with_title = bool(target and target.doc.per_equipment and plan["source"] != "client" and (
        batch_siblings > 0 or files_svc.sibling_count(session, target.ids, exclude=plan["file_id"]) > 0))
    for offset, item in enumerate(members):
        version = plan["version_no"] + offset
        file_date, _ = files_svc.file_date_for(_name(item), item.modified)
        if target is None:
            item.planned_name = naming.unclassified_name(_name(item), project.naming_templates)
        else:
            ctx = files_svc.name_context(target, original_name=_name(item), version=version,
                                         file_date=file_date, title=plan["title"],
                                         source=FileSource(plan["source"]), with_title=with_title)
            item.planned_name = naming.build_name(ctx, project.naming_templates)
        item.error = None


def counts(session: Session, batch: ImportBatch) -> dict[str, int]:
    out: dict[str, int] = defaultdict(int)
    for state, action, result, ids, suspected in session.execute(
            select(ImportItem.state, ImportItem.action, ImportItem.result,
                   ImportItem.deliverable_ids, ImportItem.suspected_equipment)
            .where(ImportItem.batch_id == batch.id)):
        out["total"] += 1
        out[f"state_{state}"] += 1
        out[f"action_{action}"] += 1
        out[f"result_{result}"] += 1
        if action == "import" and not ids:
            out["unclassified"] += 1
        if suspected:
            out["suspected"] += 1
    return dict(out)


@handler(SCAN)
def scan_job(ctx: Context) -> dict:
    s = ctx.session
    batch = s.get(ImportBatch, ctx.job.payload["batch_id"])
    try:
        total = scan(s, batch, ctx)
        analyze(s, batch)
    except Invalid as exc:
        s.rollback()
        batch.status, batch.error = "failed", exc.message
        s.commit()
        raise
    batch.status = "ready"
    s.commit()
    return {"files": total}


# ---------------------------------------------------------------- 预演报告


HEAD = {
    "zh": ("序号", "原路径", "大小 KB", "修改日期", "状态", "重复的是第几行", "识别位置", "置信度", "识别依据",
           "疑似新设备", "组内第几版", "新文件名", "处理（导入 / 跳过）", "种类编号", "设备编号", "模块",
           "来源（供应商 / 我方 / 客户）", "标题"),
    "en": ("ID", "Original path", "Size KB", "Modified", "State", "Duplicate of", "Recognised location",
           "Confidence", "Because", "Suspected new equipment", "Version in group", "New file name",
           "Action (import / skip)", "Document no.", "Equipment tag", "Module",
           "Source (vendor / ours / client)", "Title"),
}
STATE_TEXT = {"zh": {"new": "新文件", "duplicate": "本批重复", "existing": "平台已有", "old_version": "旧版",
                     "unreadable": "读不了"},
              "en": {"new": "New", "duplicate": "Duplicate", "existing": "Already in platform",
                     "old_version": "Older version", "unreadable": "Unreadable"}}
LEVEL_TEXT = {"zh": {"high": "高", "medium": "中", "low": "低"}, "en": {"high": "High", "medium": "Medium", "low": "Low"}}
ACTION_TEXT = {"zh": {"import": "导入", "skip": "跳过"}, "en": {"import": "import", "skip": "skip"}}
SOURCE_TEXT = {"zh": {"vendor": "供应商", "ours": "我方", "client": "客户"},
               "en": {"vendor": "vendor", "ours": "ours", "client": "client"}}
EDIT_FROM = 12  # 第 13 列起是可改的（0 起算）
FILL = PatternFill("solid", fgColor="FFF2CC")


def _location(session: Session, ids: list[int]) -> tuple[str, str, str]:
    """应交项 → (种类编号, 设备编号（几台用 & 连，没编号的写名称）, 模块)。"""
    if not ids:
        return "", "", ""
    ds = [session.get(Deliverable, i) for i in ids]
    doc = session.get(DocumentType, ds[0].document_type_id)
    tags = []
    for d in ds:
        if d.equipment_id:
            e = session.get(Equipment, d.equipment_id)
            tags.append(e.tag or e.name)
    module = session.get(Module, ds[0].module_id).name if ds[0].module_id else ""
    return doc.submittal_no, "&".join(tags), module


def _label(session: Session, ids: list[int]) -> str:
    """应交项 → "1.2 Datasheet / 72 Mixer & 73 Dryer"（按模块的写模块名）。"""
    if not ids:
        return ""
    ds = [session.get(Deliverable, i) for i in ids]
    doc = session.get(DocumentType, ds[0].document_type_id)
    where = []
    for d in ds:
        if d.equipment_id:
            e = session.get(Equipment, d.equipment_id)
            where.append(" ".join(x for x in (e.tag, e.name) if x))
        elif d.module_id:
            where.append(session.get(Module, d.module_id).name)
    head = f"{doc.submittal_no} {doc.name}"
    return f"{head} / {' & '.join(where)}" if where else head


def export_report(session: Session, project_id: int, batch_id: int, lang: str = "zh") -> bytes:
    batch = get_batch(session, project_id, batch_id)
    head = HEAD[lang]
    wb = Workbook()
    ws = wb.active
    ws.title = "items"
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for c in range(EDIT_FROM, len(head)):
        ws.cell(1, c + 1).fill = FILL
    labels: dict[tuple[int, ...], str] = {}
    for item in session.scalars(select(ImportItem).where(ImportItem.batch_id == batch.id)
                                .order_by(ImportItem.rel_path)):
        rec_ids = tuple(item.rec_deliverable_ids or [])
        if rec_ids and rec_ids not in labels:
            labels[rec_ids] = _label(session, list(rec_ids))
        no, tags, module = _location(session, item.deliverable_ids or [])
        reasons = "; ".join(f"{r['kind']} {r['text']}".strip() for r in item.rec_reasons or []
                            if r["kind"] != "not_delivered")
        ws.append([
            item.id, item.rel_path, round(item.size / 1024, 1), item.modified,
            STATE_TEXT[lang].get(item.state, item.state), item.duplicate_of,
            labels.get(rec_ids, ""), LEVEL_TEXT[lang].get(item.rec_level or "", ""), reasons,
            item.suspected_equipment, item.version_rank, item.planned_name or item.error,
            ACTION_TEXT[lang][item.action], no, tags, module,
            SOURCE_TEXT[lang].get(item.source or "", ""), item.title])
        for c in range(EDIT_FROM, len(head)):
            ws.cell(ws.max_row, c + 1).fill = FILL
    widths = (8, 50, 9, 11, 10, 8, 32, 8, 36, 14, 8, 50, 12, 10, 14, 12, 12, 24)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    summary = wb.create_sheet("summary")
    for k, v in sorted((batch.counts or {}).items()):
        summary.append([k, v])
    meta = wb.create_sheet("meta")
    meta.append(["docbase-import", batch.project_id, batch.id])
    meta.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _resolve(session: Session, project_id: int, no: str, tags: str, module: str
             ) -> tuple[list[int] | None, str | None]:
    """报告里写的位置 → 应交项 id；空编号 = 待归类。返回 (ids, 错误)。"""
    if not no:
        return [], None
    doc = session.scalar(select(DocumentType).where(DocumentType.project_id == project_id,
                                                    DocumentType.submittal_no == no, DocumentType.active))
    if doc is None:
        return None, f"种类编号 {no} 不存在"
    live = select(Deliverable).where(Deliverable.document_type_id == doc.id, Deliverable.obsolete.is_(False))
    if doc.per_equipment:
        if not tags:
            return None, f"{no} 是按设备的，要写设备编号"
        ids = []
        for t in [x.strip() for x in tags.split("&") if x.strip()]:
            e = session.scalar(select(Equipment).where(Equipment.project_id == project_id,
                                                       (Equipment.tag == t) | ((Equipment.tag.is_(None)) & (Equipment.name == t))))
            d = session.scalar(live.where(Deliverable.equipment_id == e.id)) if e else None
            if d is None:
                return None, f"{no} 下没有设备 {t}"
            ids.append(d.id)
        return ids, None
    if doc.per_module:
        m = session.scalar(select(Module).where(Module.project_id == project_id, Module.name == module))
        d = session.scalar(live.where(Deliverable.module_id == m.id)) if m and module else None
        return ([d.id], None) if d else (None, f"{no} 下没有模块 {module or '（空）'}")
    d = session.scalar(live.where(Deliverable.equipment_id.is_(None), Deliverable.module_id.is_(None)))
    return ([d.id], None) if d else (None, f"{no} 没有应交项")


def parse_report(session: Session, project_id: int, batch_id: int, data: bytes) -> dict[str, Any]:
    """读改过的报告：列出每行改了什么（处理、位置、来源、标题），不写入。"""
    batch = get_batch(session, project_id, batch_id)
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # noqa: BLE001  不是 Excel
        raise Invalid("import_report.bad_file", "不是 Excel 文件，或文件坏了") from exc
    if "meta" not in wb.sheetnames or wb["meta"].cell(1, 1).value != "docbase-import" or \
            wb["meta"].cell(1, 3).value != batch.id:
        raise Invalid("import_report.wrong_batch", "这不是这个导入批次导出的报告")
    ws = wb["items"]
    items = {i.id: i for i in session.scalars(select(ImportItem).where(ImportItem.batch_id == batch.id))}
    actions = {v: k for lang in ACTION_TEXT.values() for k, v in lang.items()}
    sources = {v: k for lang in SOURCE_TEXT.values() for k, v in lang.items()}
    changes, errors = [], []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, len(HEAD["zh"]) + 1)]
        if row[0] is None:
            continue
        item = items.get(row[0])
        if item is None:
            errors.append({"row": r, "message": f"第 {r} 行：序号 {row[0]} 不在这个批次里"})
            continue
        new: dict[str, Any] = {}
        action = actions.get(_text(row[12]).casefold(), actions.get(_text(row[12])))
        if action is None:
            errors.append({"row": r, "message": f"第 {r} 行：处理只能填导入 / 跳过"})
            continue
        new["action"] = action
        ids, err = _resolve(session, batch.project_id, _text(row[13]), _text(row[14]), _text(row[15]))
        if err:
            errors.append({"row": r, "message": f"第 {r} 行：{err}"})
            continue
        new["deliverable_ids"] = ids
        src = _text(row[16])
        if src and src not in sources and src.casefold() not in sources:
            errors.append({"row": r, "message": f"第 {r} 行：来源只能填供应商 / 我方 / 客户"})
            continue
        new["source"] = sources.get(src, sources.get(src.casefold())) if src else None
        new["title"] = _text(row[17]) or None
        diff = {k: [getattr(item, k), v] for k, v in new.items()
                if (sorted(getattr(item, k) or []) if k == "deliverable_ids" else getattr(item, k))
                != (sorted(v) if k == "deliverable_ids" else v)}
        if diff:
            changes.append({"item_id": item.id, "row": r, "path": item.rel_path, "changes": diff})
    return {"ok": not errors, "changes": changes, "errors": errors}


def apply_report(session: Session, project_id: int, batch_id: int, changes: list[dict[str, Any]],
                 actor: Actor) -> dict[str, int]:
    batch = get_batch(session, project_id, batch_id)
    if batch.status != "ready":
        raise Conflict("import.not_ready", "只有预演完成、还没执行的批次能改")
    for ch in changes:
        item = session.get(ImportItem, ch["item_id"])
        for field, (_, new) in ch["changes"].items():
            setattr(item, field, new)
        if item.action == "import" and item.state in ("duplicate", "existing"):
            item.state = "new"  # 用户坚持要导：执行时还会再查一次重复
    session.flush()
    analyze(session, batch, recognize=False)
    record_summary(session, project_id=project_id, entity_type="import_batches", entity_id=str(batch.id),
                   action=ChangeAction.UPDATE, changes={"report_rows": [0, len(changes)]}, actor=actor,
                   reason="导回改过的预演报告")
    return {"updated": len(changes)}


# ---------------------------------------------------------------- 执行


def start_execute(session: Session, project_id: int, batch_id: int, actor: Actor) -> ImportBatch:
    get_writable_project(session, project_id)
    batch = get_batch(session, project_id, batch_id)
    if batch.status != "ready":
        raise Conflict("import.not_ready", "只有预演完成的批次能执行")
    get_storage().check()
    batch.status = "executing"
    _queue(session, batch, EXECUTE, actor)
    record_summary(session, project_id=project_id, entity_type="import_batches", entity_id=str(batch.id),
                   action=ChangeAction.UPDATE, changes={"status": ["ready", "executing"]}, actor=actor)
    return batch


@handler(EXECUTE)
def execute_job(ctx: Context) -> dict:
    s = ctx.session
    batch = s.get(ImportBatch, ctx.job.payload["batch_id"])
    root = Path(batch.source_dir)
    storage = get_storage()
    actor = Actor(name=f"导入 #{batch.id}（{batch.created_by_name or ''}）")
    todo = s.scalars(select(ImportItem).where(ImportItem.batch_id == batch.id, ImportItem.action == "import",
                                              ImportItem.result == "pending")
                     .order_by(ImportItem.group_key, ImportItem.version_rank, ImportItem.id)).all()
    done_before = s.scalar(select(func.count()).select_from(ImportItem).where(
        ImportItem.batch_id == batch.id, ImportItem.action == "import", ImportItem.result != "pending")) or 0
    total = done_before + len(todo)
    for n, item in enumerate(todo, 1):
        item_id = item.id
        try:
            with (root / item.rel_path).open("rb") as f:
                staging_id, sha, size = storage.stage(f, max_bytes=files_svc.MAX_UPLOAD_BYTES)
            if sha != item.sha256:
                storage.discard(staging_id)
                raise Invalid("import.changed", "原文件在扫描后被改过，请重新扫描")
            overrides = {k: v for k, v in (("source", item.source), ("title", item.title)) if v}
            result = files_svc.plan(s, batch.project_id, item.deliverable_ids, original_name=_name(item),
                                    sha256=sha, modified=item.modified, overrides=overrides)
            if result["duplicate"]:
                storage.discard(staging_id)
                item.result, item.error = "skipped", "平台已有内容完全相同的文件"
            else:
                stored = files_svc.store_staged(s, batch.project_id, storage, staging_id=staging_id,
                                                sha256=sha, size=size, original_name=_name(item),
                                                result=result, actor=actor)
                item.result, item.file_id, item.version_id, item.error = \
                    "done", stored["file_id"], stored["version_id"], None
            s.commit()
        except Exception as exc:  # noqa: BLE001  逐条记失败原因，接着做下一个
            s.rollback()
            item = s.get(ImportItem, item_id)
            item.result = "failed"
            item.error = getattr(exc, "message", None) or f"{type(exc).__name__}: {exc}"
            s.commit()
        ctx.heartbeat(progress=done_before + n, total=total)
    batch = s.get(ImportBatch, batch.id)
    batch.status, batch.executed_at = "done", datetime.now(UTC)
    batch.counts = counts(s, batch)
    emit(s, "import.executed", batch, {"done": batch.counts.get("result_done", 0)})
    s.commit()
    return {"done": batch.counts.get("result_done", 0), "failed": batch.counts.get("result_failed", 0)}


# ---------------------------------------------------------------- 撤销


def start_undo(session: Session, project_id: int, batch_id: int, actor: Actor) -> ImportBatch:
    get_writable_project(session, project_id)
    batch = get_batch(session, project_id, batch_id)
    if batch.status != "done":
        raise Conflict("import.not_done", "只有执行完的批次能撤销")
    batch.status = "undoing"
    _queue(session, batch, UNDO, actor)
    return batch


@handler(UNDO)
def undo_job(ctx: Context) -> dict:
    """删除这一批建的版本（没别的版本了就连文件一起删）；已提交给客户的、导入后又被人升过版的不撤。"""
    s = ctx.session
    batch = s.get(ImportBatch, ctx.job.payload["batch_id"])
    actor = Actor(name=f"撤销导入 #{batch.id}", id=ctx.job.created_by)
    by_file: dict[int, list[ImportItem]] = defaultdict(list)
    for item in s.scalars(select(ImportItem).where(ImportItem.batch_id == batch.id,
                                                   ImportItem.result == "done")):
        by_file[item.file_id].append(item)
    keep = get_settings().keep_versions
    for n, (file_id, items) in enumerate(sorted(by_file.items()), 1):
        mine = {i.version_id for i in items}
        versions = s.scalars(select(Version).where(Version.file_id == file_id)
                             .order_by(Version.version_no)).all()
        protected = files_svc.protected_version_ids(s, list(mine))
        lowest = min((v.version_no for v in versions if v.id in mine), default=0)
        later_others = [v for v in versions if v.id not in mine and v.version_no > lowest]
        if protected or later_others:
            reason = "已提交给客户（在传送单或阶段基线里）" if protected else "导入后又被人升过版"
            for i in items:
                i.result, i.error = "kept", f"没有撤销：{reason}"
            s.commit()
            ctx.heartbeat(progress=n, total=len(by_file))
            continue
        file = s.get(File, file_id)
        for v in versions:
            if v.id in mine:
                record(s, v, ChangeAction.DELETE, snapshot(v), None, actor, project_id=batch.project_id,
                       reason=f"撤销导入 #{batch.id}")
                s.delete(v)
        s.flush()
        remaining = [v for v in versions if v.id not in mine]
        if remaining:
            before = snapshot(file)
            file.current_version_no = max(v.version_no for v in remaining)
            # 因为这批导入被挤进回收区的旧版本，放回来
            for v in sorted(remaining, key=lambda v: -v.version_no)[:keep]:
                if v.recycled_at is not None and v.purged_at is None:
                    v.recycled_at = None
            record(s, file, ChangeAction.UPDATE, before, snapshot(file), actor, reason=f"撤销导入 #{batch.id}")
        else:
            record(s, file, ChangeAction.DELETE, snapshot(file), None, actor, reason=f"撤销导入 #{batch.id}")
            for link in s.scalars(select(FileDeliverable).where(FileDeliverable.file_id == file_id)):
                s.delete(link)
            for i in items:
                i.file_id = None
            s.flush()
            s.delete(file)
        for i in items:
            i.result, i.version_id = "undone", None
        s.commit()
        ctx.heartbeat(progress=n, total=len(by_file))
    batch = s.get(ImportBatch, batch.id)
    batch.status, batch.undone_at = "undone", datetime.now(UTC)
    batch.counts = counts(s, batch)
    emit(s, "import.undone", batch, {"kept": batch.counts.get("result_kept", 0)})
    s.commit()
    return {"undone": batch.counts.get("result_undone", 0), "kept": batch.counts.get("result_kept", 0)}


# ---------------------------------------------------------------- 明细列表


def list_items(session: Session, project_id: int, batch_id: int, *, state: str | None = None,
               result: str | None = None, unclassified: bool | None = None, suspected: bool | None = None,
               offset: int = 0, limit: int = 100) -> dict[str, Any]:
    batch = get_batch(session, project_id, batch_id)
    stmt = select(ImportItem).where(ImportItem.batch_id == batch.id)
    if state:
        stmt = stmt.where(ImportItem.state == state)
    if result:
        stmt = stmt.where(ImportItem.result == result)
    if suspected:
        stmt = stmt.where(ImportItem.suspected_equipment.is_not(None))
    rows = session.scalars(stmt.order_by(ImportItem.rel_path)).all()
    if unclassified:
        rows = [r for r in rows if r.action == "import" and not r.deliverable_ids]
    page = rows[offset:offset + limit]
    out = []
    for r in page:
        out.append({"id": r.id, "path": r.rel_path, "size": r.size, "modified": r.modified,
                    "state": r.state, "duplicate_of": r.duplicate_of, "level": r.rec_level,
                    "reasons": r.rec_reasons, "suspected_equipment": r.suspected_equipment,
                    "action": r.action, "deliverable_ids": r.deliverable_ids,
                    "location": _label(session, r.deliverable_ids or []),
                    "manual": sorted(r.deliverable_ids or []) != sorted(r.rec_deliverable_ids or []),
                    "source": r.source, "title": r.title, "version_rank": r.version_rank,
                    "planned_name": r.planned_name, "result": r.result, "file_id": r.file_id,
                    "error": r.error})
    return {"total": len(rows), "offset": offset, "limit": limit, "items": out}
