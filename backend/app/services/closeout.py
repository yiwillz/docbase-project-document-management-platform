"""闭环检查与里程碑关闭（PRD F7.7）。

一个里程碑能关闭：每个应交项都已批准（带意见批准的要拿到最终批准）或不适用，
并且有凭证（批准有客户回复；不适用有依据）；客户新增要求都已确认接受与否。
内部要求的问题、文件种类没写要求依据，只提示不挡关闭。
批准之后又传了新版本（我方 / 供应商文件）的也挡关闭：客户批准的不是最新那一版。

关闭只是记录（日期、操作人、依据），不锁定；关闭后有变动会提示，可以重新打开。
"""

import io
from datetime import UTC, date, datetime
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    Acceptance,
    ChangeAction,
    ChangeLog,
    DeliverableStatus,
    DocumentType,
    DocumentTypeOrigin,
    Milestone,
)
from app.services import mdl_export as mx
from app.services.crud import get_in_project
from app.services.deliverable_query import Filters, Item, load_items, stats
from app.services.errors import Conflict, Invalid
from app.services.files import versions_after_approval
from app.services.history import Actor, emit, record, snapshot
from app.services.projects import get_project, get_writable_project
from app.services.transmittals import latest_items, review_info

S = DeliverableStatus
# 问题代码：挡关闭的 / 只提示的
BLOCKING = ("not_done", "not_final", "no_evidence", "requirement_unconfirmed",
            "new_version_after_approval")
WARNING = ("no_requirement_basis", "internal_open")

ISSUE_TEXT = {
    "en": {"not_done": "Not approved", "not_final": "Final approval pending (approved with comments)",
           "no_evidence": "No evidence (no client reply / no basis)",
           "requirement_unconfirmed": "Client requirement not yet accepted or declined",
           "new_version_after_approval": "New version uploaded after approval",
           "no_requirement_basis": "Requirement basis missing",
           "internal_open": "Internal requirement open"},
    "zh": {"not_done": "未批准", "not_final": "带意见批准，尚未拿到最终批准",
           "no_evidence": "无凭证（没有客户回复 / 没写依据）",
           "requirement_unconfirmed": "客户新增要求尚未确认接受与否",
           "new_version_after_approval": "批准后又传了新版本",
           "no_requirement_basis": "没写要求依据", "internal_open": "内部要求未完成"},
}
ORIGIN_TEXT = {"en": {"contract": "Contract", "client": "Client request", "internal": "Internal"},
               "zh": {"contract": "合同规定", "client": "客户要求", "internal": "内部要求"}}


def _item_issues(it: Item, acceptance: dict[int, Acceptance | None], newer: set[int]) -> list[str]:
    issues = []
    done = it.status in (S.APPROVED.value, S.NOT_APPLICABLE.value)
    if not done:
        issues.append("not_done")
    if it.status == S.APPROVED.value and it.review_final is False:
        issues.append("not_final")
    if it.status == S.APPROVED.value and it.id in newer:
        issues.append("new_version_after_approval")
    if it.no_evidence:
        issues.append("no_evidence")
    if acceptance.get(it.document_type_id) is Acceptance.PENDING:
        issues.append("requirement_unconfirmed")
    if it.origin == DocumentTypeOrigin.INTERNAL.value and issues:
        return ["internal_open"]  # 内部要求：只提示
    return issues


def check(session: Session, project_id: int, milestone_id: int, *, today: date | None = None
          ) -> dict[str, Any]:
    get_project(session, project_id)
    ms = get_in_project(session, Milestone, milestone_id, project_id, "milestone")
    items = load_items(session, project_id, Filters(milestone_ids=[ms.id]), today=today)
    docs = list(session.scalars(select(DocumentType).where(
        DocumentType.milestone_id == ms.id, DocumentType.active).order_by(DocumentType.sort_order)))
    acceptance = {d.id: d.acceptance for d in docs}
    reviews = review_info(session, project_id, today or date.today())
    approved_ids = [it.id for it in items if it.status == S.APPROVED.value]
    latest = latest_items(session, project_id, approved_ids)
    newer = versions_after_approval(session, project_id, {
        did: x.item for did, x in latest.items() if x.item.review_code is not None})

    rows, counts = [], dict.fromkeys(BLOCKING + WARNING, 0)
    for it in items:
        issues = _item_issues(it, acceptance, newer)
        for code in issues:
            counts[code] += 1
        review = reviews.get(it.id, {})
        rows.append({**it.out(), "revision": review.get("revision"), "issues": issues})
    documents = []
    for d in docs:
        if d.origin is not DocumentTypeOrigin.INTERNAL and not (d.basis or "").strip():
            counts["no_requirement_basis"] += 1
            documents.append({"document_type_id": d.id, "submittal_no": d.submittal_no,
                              "name": d.name, "origin": d.origin.value,
                              "issues": ["no_requirement_basis"]})
    blocking = sum(1 for r in rows if any(i in BLOCKING for i in r["issues"]))
    changed_after_close = ms.closed_at is not None and _changed_after_close(session, ms, items)
    return {
        "milestone": {"id": ms.id, "code": ms.code, "name": ms.name, "due_date": ms.due_date,
                      "closed": ms.closed_at is not None, "closed_at": ms.closed_at,
                      "closed_date": ms.closed_date, "closed_by": ms.closed_by_name,
                      "close_basis": ms.close_basis, "changed_after_close": changed_after_close},
        "closable": blocking == 0, "blocking": blocking,
        "warnings": sum(1 for r in rows if any(i in WARNING for i in r["issues"])) + len(documents),
        "counts": counts, "stats": stats(items), "documents": documents, "items": rows,
    }


def _changed_after_close(session: Session, ms: Milestone, items: list[Item]) -> bool:
    """关闭之后，这个里程碑的应交项有没有改动（按变更记录的先后判断）。"""
    closed_log = session.scalar(select(func.max(ChangeLog.id)).where(
        ChangeLog.entity_type == "milestones", ChangeLog.entity_id == str(ms.id)))
    if closed_log is None or not items:
        return False
    return session.scalar(select(func.count()).select_from(ChangeLog).where(
        ChangeLog.entity_type == "deliverables", ChangeLog.id > closed_log,
        ChangeLog.entity_id.in_([str(it.id) for it in items]))) > 0


def close(session: Session, project_id: int, milestone_id: int, data: dict[str, Any],
          actor: Actor) -> Milestone:
    """记录关闭：日期（真实日期，默认今天）、依据（如付款申请号）。还有挡关闭的问题时，
    要 force 并写原因（如客户书面同意带尾项关闭）。"""
    get_writable_project(session, project_id)
    ms = get_in_project(session, Milestone, milestone_id, project_id, "milestone")
    if ms.closed_at is not None:
        raise Conflict("milestone.closed", f"里程碑 {ms.code} 已关闭", id=ms.id)
    result = check(session, project_id, ms.id)
    reason = (data.get("reason") or "").strip() or None
    if not result["closable"]:
        if not data.get("force"):
            raise Conflict("milestone.not_closable", f"里程碑 {ms.code} 还有 {result['blocking']} 项未闭环",
                           blocking=result["blocking"])
        if reason is None:
            raise Invalid("milestone.force_needs_reason", "带未闭环项关闭必须写原因")
    before = snapshot(ms)
    ms.closed_at = datetime.now(UTC)
    ms.closed_date = data.get("closed_date") or date.today()
    ms.closed_by_name = actor.name
    basis = (data.get("basis") or "").strip() or None
    if not result["closable"]:
        basis = f"{basis or ''}（带 {result['blocking']} 项未闭环关闭：{reason}）".lstrip()
    ms.close_basis = basis
    record(session, ms, ChangeAction.UPDATE, before, snapshot(ms), actor, reason=reason,
           project_id=project_id)
    emit(session, "milestone.closed", ms, {"code": ms.code, "blocking": result["blocking"]},
         project_id=project_id)
    return ms


def reopen(session: Session, project_id: int, milestone_id: int, actor: Actor, *,
           reason: str | None) -> Milestone:
    get_writable_project(session, project_id)
    ms = get_in_project(session, Milestone, milestone_id, project_id, "milestone")
    if ms.closed_at is None:
        raise Conflict("milestone.not_closed", f"里程碑 {ms.code} 没有关闭", id=ms.id)
    if not (reason or "").strip():
        raise Invalid("milestone.reopen_needs_reason", "重新打开必须写原因")
    before = snapshot(ms)
    ms.closed_at = ms.closed_date = ms.closed_by_name = ms.close_basis = None
    record(session, ms, ChangeAction.UPDATE, before, snapshot(ms), actor, reason=reason,
           project_id=project_id)
    emit(session, "milestone.reopened", ms, {"code": ms.code}, project_id=project_id)
    return ms


# ---------------------------------------------------------------- 关闭清单导出

TEXT = {
    "en": {"title": "MILESTONE CLOSEOUT LIST", "project": "Project", "milestone": "Milestone",
           "client": "Client", "result": "Result", "closable": "Ready to close",
           "open": "{n} item(s) open", "closed": "Closed {date} ({basis})", "as_of": "As of",
           "columns": ("No.", "Document Title", "Module / Equipment", "Requirement", "Status",
                       "Code", "Reply Date", "Reply Ref", "Transmittal", "Rev.", "Issues"),
           "sign": ("Prepared by (Contractor)", "Accepted by (Client)"),
           "sign_fields": ("Name", "Signature", "Date"), "sheet": "Closeout",
           "footer": "Generated by {p}  ·  {code} {ms}  ·  {at}  ·  System-generated document – do not edit manually",
           "page": "Page {n} of {m}"},
    "zh": {"title": "里程碑关闭清单", "project": "项目", "milestone": "里程碑", "client": "客户",
           "result": "结论", "closable": "可以关闭", "open": "还有 {n} 项未闭环",
           "closed": "已于 {date} 关闭（{basis}）", "as_of": "统计日期",
           "columns": ("编号", "文件名称", "模块 / 设备", "要求类别", "状态", "审核代码", "回复日期",
                       "回函号", "传送单", "版次", "问题"),
           "sign": ("编制（总包）", "确认（客户）"), "sign_fields": ("姓名", "签名", "日期"),
           "sheet": "关闭清单",
           "footer": "{p} 系统生成  ·  {code} {ms}  ·  {at}  ·  系统生成文件，请勿手工修改",
           "page": "第 {n} 页 / 共 {m} 页"},
}


def _rows(result: dict[str, Any], lang: str) -> list[tuple]:
    labels, issues_t, origin_t = mx.TEXT[lang]["status"], ISSUE_TEXT[lang], ORIGIN_TEXT[lang]
    rows = []
    for r in result["items"]:
        key = "O" if r["overdue"] else mx.KEY_OF[S(r["status"])]
        tag = r["equipment_tag"]
        if r["equipment_id"]:
            where = f"{tag} {r['equipment_name']}" if tag else r["equipment_name"]
        else:
            where = r["module_name"] or ""
        no = f"{r['submittal_no']}.{tag}" if tag else r["submittal_no"]
        rows.append((no, r["document_type_name"], where, origin_t[r["origin"]], labels[key],
                     r["review_code"], r["reply_date"], r["reply_ref"], r["transmittal_no"],
                     r["revision"], "; ".join(issues_t[i] for i in r["issues"]), key,
                     bool(set(r["issues"]) & set(BLOCKING))))
    return rows


def _verdict(result: dict[str, Any], t: dict[str, str]) -> str:
    ms = result["milestone"]
    if ms["closed"]:
        return t["closed"].format(date=ms["closed_date"], basis=ms["close_basis"] or "")
    return t["closable"] if result["closable"] else t["open"].format(n=result["blocking"])


def export(session: Session, project_id: int, milestone_id: int, *, lang: str = "en",
           fmt: str = "xlsx", today: date | None = None) -> bytes:
    project = get_project(session, project_id)
    result = check(session, project_id, milestone_id, today=today)
    t = TEXT[lang]
    ms = result["milestone"]
    ctx = {"t": t, "lang": lang, "project": project, "result": result, "rows": _rows(result, lang),
           "verdict": _verdict(result, t), "today": today or date.today(),
           "footer": t["footer"].format(p=mx.PLATFORM, code=project.code, ms=ms["code"],
                                        at=f"{datetime.now():%Y-%m-%d %H:%M}")}
    return _excel(ctx) if fmt == "xlsx" else _pdf(ctx)


def _head_rows(ctx: dict[str, Any]) -> list[tuple[str, Any]]:
    t, p, ms = ctx["t"], ctx["project"], ctx["result"]["milestone"]
    return [(t["project"], f"{p.name} ({p.code})"), (t["client"], p.client_name or ""),
            (t["milestone"], f"{ms['code']} {ms['name'] or ''}"), (t["result"], ctx["verdict"]),
            (t["as_of"], f"{ctx['today']:%Y-%m-%d}")]


def _excel(ctx: dict[str, Any]) -> bytes:
    t = ctx["t"]
    wb = Workbook()
    ws = wb.active
    ws.title = t["sheet"]
    ws["A1"] = t["title"]
    ws["A1"].font = Font(size=16, bold=True, color=mx.NAVY)
    for r, (k, v) in enumerate(_head_rows(ctx), start=2):
        ws.cell(r, 1, k).font = Font(bold=True, color="595959")
        ws.cell(r, 2, v)
    ws.cell(5, 2).font = Font(bold=True, color="006100" if ctx["result"]["closable"] else "9C0006")
    header = 8
    mx._header_cells(ws, header, t["columns"], (10, 36, 26, 14, 18, 8, 12, 24, 20, 8, 44))
    for r, row in enumerate(ctx["rows"], start=header + 1):
        for col, v in enumerate(row[:11], start=1):
            c = ws.cell(r, col, v)
            c.border = mx.BOX
            c.alignment = Alignment(vertical="center", wrap_text=col in (2, 11))
            if isinstance(v, date):
                c.number_format = "yyyy-mm-dd"
        fill, font = mx.COLORS[row[11]]
        sc = ws.cell(r, 5)
        sc.fill, sc.font = PatternFill("solid", fgColor=fill), Font(bold=True, color=font)
        if row[12]:
            ws.cell(r, 11).font = Font(color="9C0006")
    last = header + len(ctx["rows"])
    ws.auto_filter.ref = f"A{header}:K{last}"
    ws.freeze_panes = ws.cell(header + 1, 1)
    r = last + 3
    for col, role in ((2, t["sign"][0]), (8, t["sign"][1])):
        ws.cell(r, col, role).font = Font(bold=True, color=mx.NAVY)
        for j, label in enumerate(t["sign_fields"], start=1):
            ws.cell(r + j, col, f"{label}:").font = Font(color="595959")
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header}:{header}"
    ws.oddFooter.left.text, ws.oddFooter.left.size = ctx["footer"], 7
    wb.properties.creator = f"{mx.PLATFORM} (system generated)"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _pdf(ctx: dict[str, Any]) -> bytes:
    t = ctx["t"]
    regular, bold = mx._fonts(ctx["lang"])
    navy, grid = colors.HexColor("#" + mx.NAVY), colors.HexColor("#" + mx.GRID)
    small = ParagraphStyle("s", fontName=regular, fontSize=7, leading=8.5)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=t["title"],
                            author=f"{mx.PLATFORM} (system generated)")
    head = Table([[k, v] for k, v in _head_rows(ctx)], colWidths=[25 * mm, 248 * mm])
    head.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), regular, 8.5), ("FONT", (0, 0), (0, -1), bold, 8.5),
        ("FONT", (1, 3), (1, 3), bold, 8.5),
        ("TEXTCOLOR", (1, 3), (1, 3), colors.HexColor("#006100" if ctx["result"]["closable"] else "#9C0006")),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#595959")),
        ("LINEBELOW", (0, -1), (-1, -1), 1.2, navy)]))
    data: list[list[Any]] = [list(t["columns"])]
    style = [("FONT", (0, 0), (-1, 0), bold, 7), ("BACKGROUND", (0, 0), (-1, 0), navy),
             ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONT", (0, 1), (-1, -1), regular, 7),
             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("GRID", (0, 0), (-1, -1), 0.3, grid)]
    for row in ctx["rows"]:
        data.append([row[0], Paragraph(escape(row[1]), small), Paragraph(escape(row[2]), small), row[3],
                     row[4], row[5] or "", f"{row[6]:%Y-%m-%d}" if row[6] else "",
                     Paragraph(escape(row[7] or ""), small), row[8] or "", row[9] or "",
                     Paragraph(f"<font color='#9C0006'>{escape(row[10])}</font>" if row[12]
                               else escape(row[10]), small)])
        i = len(data) - 1
        fill, font = mx.COLORS[row[11]]
        style += [("BACKGROUND", (4, i), (4, i), colors.HexColor("#" + fill)),
                  ("TEXTCOLOR", (4, i), (4, i), colors.HexColor("#" + font))]
    table = Table(data, colWidths=[w * mm for w in (16, 40, 38, 18, 24, 10, 17, 26, 31, 10, 43)],
                  repeatRows=1)
    table.setStyle(TableStyle(style))
    f_name, f_sign, f_date = t["sign_fields"]
    sign = Table([list(t["sign"]), [f"{f_name}:"] * 2, [f"{f_sign}:"] * 2, [f"{f_date}:"] * 2],
                 colWidths=[136 * mm] * 2, rowHeights=[6 * mm, 8 * mm, 10 * mm, 8 * mm])
    sign.setStyle(TableStyle(
        [("FONT", (0, 0), (-1, -1), regular, 8), ("FONT", (0, 0), (-1, 0), bold, 8),
         ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
         ("VALIGN", (0, 1), (-1, -1), "BOTTOM")]
        + [("BOX", (i, 0), (i, -1), 0.5, grid) for i in range(2)]))
    story = [Paragraph(f"<font size=16 color='#{mx.NAVY}'><b>{escape(t['title'])}</b></font>",
                       ParagraphStyle("t", parent=small, fontName=bold, leading=20, spaceAfter=3 * mm)),
             head, Spacer(1, 5 * mm), table, Spacer(1, 6 * mm), KeepTogether([sign])]
    doc.build(story, canvasmaker=mx.numbered_canvas(regular, ctx["footer"], t["page"]))
    return buffer.getvalue()
