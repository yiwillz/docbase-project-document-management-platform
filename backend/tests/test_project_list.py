"""第 17C 步：项目列表的完成率 / 缺件 / 逾期（PRD F1.2）、变更记录导出（PRD F10.2）。"""

import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app.models import Deliverable, DeliverableStatus

CSRF = {"X-Docbase": "1"}


def _project(api_client, code):
    return api_client.post("/api/projects", json={"code": code, "name": f"{code} demo"}, headers=CSRF).json()


def _url(project, path):
    return f"/api/projects/{project['id']}/{path}"


def test_project_list_stats_match_the_matrix(api_client, session):
    p = _project(api_client, "PL-1")
    empty = _project(api_client, "PL-2")
    past = (date.today() - timedelta(days=3)).isoformat()
    m = api_client.post(_url(p, "milestones"), json={"name": "Design", "due_date": past}, headers=CSRF).json()
    api_client.post(_url(p, "document-types/bulk"), headers=CSRF, json={
        "milestone_id": m["id"], "rows": [{"name": n} for n in ("A", "B", "C", "D", "E")]})
    items = api_client.get(_url(p, "deliverables")).json()["items"]
    status = {it["submittal_no"]: it["id"] for it in items}
    for no, st in (("1.1", DeliverableStatus.APPROVED), ("1.2", DeliverableStatus.NOT_APPLICABLE),
                   ("1.3", DeliverableStatus.REVISE_RESUBMIT), ("1.4", DeliverableStatus.SUBMITTED)):
        session.get(Deliverable, status[no]).status = st
    session.flush()
    # 1.4（已提交）收到了文件
    r = api_client.post(_url(p, "files/preview"), data={"deliverable_ids": str(status["1.4"])}, headers=CSRF,
                        files={"file": ("D.pdf", b"d", "application/pdf")})
    api_client.post(_url(p, f"files/uploads/{r.json()['token']}/confirm"), json={}, headers=CSRF)

    rows = {r["code"]: r for r in api_client.get("/api/projects").json()}
    got = rows["PL-1"]
    # 应交 4（不含不适用），已批准 1；缺件 = 没收到文件的 1.5 + 需修改重交的 1.3（1.4 收到了文件不算）；
    # 逾期 = 未提交且过了截止日 1.5
    assert (got["completion"], got["outstanding_count"], got["overdue_count"]) == (0.25, 2, 1)
    assert (rows["PL-2"]["completion"], rows["PL-2"]["outstanding_count"]) == (None, 0)  # 没有应交项
    # 和应交矩阵的汇总一样
    overall = api_client.get(_url(p, "deliverables/summary")).json()["overall"]
    assert (overall["completion"], overall["overdue"], overall["outstanding"]) == (
        got["completion"], got["overdue_count"], got["outstanding_count"])
    assert empty["id"] in {r["id"] for r in api_client.get("/api/projects").json()}


def test_change_log_export(api_client):
    p = _project(api_client, "CL-1")
    m = api_client.post(_url(p, "milestones"), json={"name": "Design"}, headers=CSRF).json()
    api_client.patch(_url(p, f"milestones/{m['id']}"), json={"name": "Detailed design", "reason": "客户改名"},
                     headers=CSRF)

    resp = api_client.get(_url(p, "change-logs/export.xlsx"), params={"entity_type": "milestones"})
    assert resp.status_code == 200 and "CL-1_ChangeLog_" in resp.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws["A1"].value == "变更记录" and "里程碑" in ws["B3"].value
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6)]
    assert all(r[2] == "里程碑" for r in rows)
    renamed = [r for r in rows if r[5] == "name" and r[4] == "修改"]
    assert renamed and (renamed[0][6], renamed[0][7], renamed[0][8]) == ("Design", "Detailed design", "客户改名")
    assert any(r[4] == "新建" for r in rows)

    en = load_workbook(io.BytesIO(api_client.get(_url(p, "change-logs/export.xlsx"),
                                                 params={"lang": "en"}).content)).active
    assert en["A1"].value == "CHANGE LOG" and en.cell(5, 1).value == "Time"


def test_change_log_export_truncates(api_client, monkeypatch):
    from app.services import change_log_export

    monkeypatch.setattr(change_log_export, "MAX_ROWS", 3)
    p = _project(api_client, "CL-2")
    for i in range(3):
        api_client.post(_url(p, "milestones"), json={"name": f"M{i}"}, headers=CSRF)
    ws = load_workbook(io.BytesIO(api_client.get(_url(p, "change-logs/export.xlsx")).content)).active
    assert "只导出了最新的 3 行" in ws["A4"].value
    assert sum(1 for _ in ws.iter_rows(min_row=6)) <= 3


# ---------------------------------------------------------------- 第 17E 步：删除空项目、变更记录的对象名称


def test_delete_empty_project(api_client, session):
    from sqlalchemy import select

    from app.models import ChangeLog, Milestone, Project

    p = _project(api_client, "DEL-1")
    m = api_client.post(_url(p, "milestones"), json={"name": "Design"}, headers=CSRF).json()
    api_client.post(_url(p, "document-types/bulk"), json={"milestone_id": m["id"], "rows": [{"name": "Spec"}]},
                    headers=CSRF)
    alpha = api_client.post(_url(p, "equipment-types"), json={"name": "Alpha"}, headers=CSRF).json()
    api_client.post(_url(p, "equipment"), json={"tag": "72", "name": "Mixer", "equipment_type_id": alpha["id"]},
                    headers=CSRF)
    assert api_client.get(_url(p, "delete-check")).json() == {"deletable": True, "blockers": {}}

    # 编号输错不删
    r = api_client.post(_url(p, "delete"), json={"confirm_code": "DEL-X"}, headers=CSRF)
    assert r.json()["detail"]["code"] == "project.delete_confirm"
    r = api_client.post(_url(p, "delete"), json={"confirm_code": "DEL-1", "reason": "建错了"}, headers=CSRF)
    assert r.status_code == 200 and r.json()["counts"]["milestones"] == 1
    assert session.get(Project, p["id"]) is None
    assert not session.scalars(select(Milestone).where(Milestone.project_id == p["id"])).all()
    # 项目下的变更记录随项目删掉，另留一条不挂项目的删除记录
    logs = session.scalars(select(ChangeLog).where(ChangeLog.entity_type == "projects",
                                                   ChangeLog.entity_id == str(p["id"]))).all()
    assert [(log.project_id, log.action.value, log.reason) for log in logs] == [(None, "delete", "建错了")]
    assert logs[0].changes["code"] == ["DEL-1", None]


def test_project_with_files_cannot_be_deleted(api_client):
    p = _project(api_client, "DEL-2")
    m = api_client.post(_url(p, "milestones"), json={"name": "Design"}, headers=CSRF).json()
    api_client.post(_url(p, "document-types/bulk"), json={"milestone_id": m["id"], "rows": [{"name": "Spec"}]},
                    headers=CSRF)
    did = api_client.get(_url(p, "deliverables")).json()["items"][0]["id"]
    r = api_client.post(_url(p, "files/preview"), data={"deliverable_ids": str(did)}, headers=CSRF,
                        files={"file": ("a.pdf", b"a", "application/pdf")})
    api_client.post(_url(p, f"files/uploads/{r.json()['token']}/confirm"), json={}, headers=CSRF)
    assert api_client.get(_url(p, "delete-check")).json() == {"deletable": False, "blockers": {"files": 1}}
    r = api_client.post(_url(p, "delete"), json={"confirm_code": "DEL-2"}, headers=CSRF)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "project.not_empty"


def test_change_log_labels_for_links_and_summaries(api_client):
    p = _project(api_client, "LB-1")
    alpha = api_client.post(_url(p, "equipment-types"), json={"name": "Alpha"}, headers=CSRF).json()
    mod = api_client.post(_url(p, "modules"), json={"name": "A"}, headers=CSRF).json()
    api_client.put(_url(p, f"modules/{mod['id']}/equipment-types"), json={"equipment_type_ids": [alpha["id"]]},
                   headers=CSRF)
    m = api_client.post(_url(p, "milestones"), json={"name": "Design"}, headers=CSRF).json()
    api_client.post(_url(p, "document-types/bulk"), json={"milestone_id": m["id"], "rows": [{"name": "Spec"}]},
                    headers=CSRF)
    items = api_client.get("/api/change-logs", params={"project_id": p["id"], "limit": 200}).json()["items"]
    link = next(i for i in items if i["entity_type"] == "module_equipment_types")
    assert link["label"] == "A / Alpha" and link["summary"] is None
    recalc = next(i for i in items if i["entity_type"] == "deliverables" and i["entity_id"].startswith("project:"))
    assert recalc["summary"] == "recalculate"
    # 导出里也写成说明，不是 project:N
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(api_client.get(_url(p, "change-logs/export.xlsx")).content)).active
    objects = {row[3] for row in ws.iter_rows(min_row=6, values_only=True)}
    assert "重算应交项" in objects and "A / Alpha" in objects
    assert not any(str(o).startswith("project:") for o in objects)
