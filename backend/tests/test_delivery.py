"""第 15 步：交付包、阶段基线、共享盘镜像、传送单记全部版本。

虚构项目 DL-1：设备种类 Alpha；72 Mixer、73 Dryer（都在 Alpha）
M1 Design
  1.1 整个种类 Spec：Area A.pdf；客户文件 Client spec.pdf
  1.2 按设备 Datasheet：72 有两份（Datasheet.pdf、Wiring.dwg）；72&73 共用一份 Shared.pdf
  1.3 按模块 Layout：A
另有一份待归类 notes.txt
"""

import io
import os
import stat
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import worker
from app.models import Deliverable, Job, TransmittalItem, TransmittalItemVersion, Version
from app.services import delivery, mirror
from app.services import files as files_svc
from app.services.storage import get_storage

USER = {"X-User": "alice"}


def url(project, path):
    return f"/api/projects/{project['id']}/{path}"


def call(api_client, method, project, path, body=None, status=200, **kw):
    resp = api_client.request(method, url(project, path), json=body, headers=USER, **kw)
    assert resp.status_code == status, resp.text
    return resp.json() if resp.headers.get("content-type", "").startswith("application/json") else resp


def store(api_client, project, ids, name, content, overrides=None):
    resp = api_client.post(url(project, "files/preview"), headers=USER,
                           data={"deliverable_ids": ",".join(map(str, ids))}, files={"file": (name, content)})
    assert resp.status_code == 200, resp.text
    return call(api_client, "POST", project, f"files/uploads/{resp.json()['token']}/confirm", overrides or {})


@pytest.fixture
def project(api_client):
    return api_client.post("/api/projects", json={"code": "DL-1", "name": "Delivery demo"}, headers=USER).json()


@pytest.fixture
def d(api_client, project, session):
    alpha = call(api_client, "POST", project, "equipment-types", {"name": "Alpha"}, 201)
    m1 = call(api_client, "POST", project, "milestones", {"name": "Design"}, 201)
    docs = call(api_client, "POST", project, "document-types/bulk", {
        "milestone_id": m1["id"], "rows": [{"name": "Spec"}, {"name": "Datasheet"}, {"name": "Layout"}]}, 201)
    eq = {}
    for tag, name in (("72", "Mixer"), ("73", "Dryer")):
        eq[tag] = call(api_client, "POST", project, "equipment",
                       {"tag": tag, "name": name, "equipment_type_id": alpha["id"]}, 201)["id"]
    call(api_client, "POST", project, "document-types/bulk-equipment", {
        "document_type_ids": [docs[1]["id"]], "all_equipment": True, "mode": "add"})
    call(api_client, "PUT", project, f"document-types/{docs[2]['id']}/modules", {"module_names": ["A"]})
    out = {"m1": m1["id"], "docs": [x["id"] for x in docs], "eq": eq}
    for row in session.scalars(select(Deliverable).where(Deliverable.project_id == project["id"],
                                                         Deliverable.obsolete.is_(False))):
        key = {eq["72"]: "72", eq["73"]: "73"}.get(row.equipment_id) or ("A" if row.module_id else "spec")
        out[key] = row.id
    return out


@pytest.fixture
def stored(api_client, project, d):
    return {
        "spec": store(api_client, project, [d["spec"]], "Area A 20260226.pdf", b"spec"),
        "client": store(api_client, project, [d["spec"]], "Client spec.pdf", b"client", {"source": "client"}),
        "ds": store(api_client, project, [d["72"]], "Datasheet 20260226.pdf", b"ds"),
        "wiring": store(api_client, project, [d["72"]], "Wiring 20260226.dwg", b"wiring"),
        "shared": store(api_client, project, [d["72"], d["73"]], "Shared 20260226.pdf", b"shared"),
        "layout": store(api_client, project, [d["A"]], "Plan 20260226.dwg", b"plan"),
    }


def paths(session, project, **kw):
    return [e.path for e in delivery.entries(session, project["id"], **kw)]


# ---------------------------------------------------------------- 目录结构


def test_equipment_layout(session, project, d, stored):
    p = paths(session, project)
    assert "DL-1/通用文件/M1 Design/1.1 Spec/1.1_Spec-Area A-20260226-V1.pdf" in p
    # 72 下有两份文件：名字里带标题
    assert "DL-1/Alpha/72 Mixer/1.2.72_Datasheet-Mixer-Datasheet-20260226-V1.pdf" in p
    assert "DL-1/Alpha/72 Mixer/1.2.72_Datasheet-Mixer-Wiring-20260226-V1.dwg" in p
    # 72&73 共用的：两台设备的文件夹里各一份
    shared = [x for x in p if "72&73_" in x]  # 72&73 上只有这一份：名字里不带标题
    assert {x.rsplit("/", 2)[1] for x in shared} == {"72 Mixer", "73 Dryer"}
    assert "DL-1/通用文件/M1 Design/1.3 Layout/A/1.3_Layout-Plan-20260226-V1.dwg" in p
    assert not any("Client spec" in x for x in p)  # 默认不放客户文件
    with_client = paths(session, project, include_client=True)
    assert "DL-1/客户文件/M1 Design/1.1 Spec/Client spec.pdf" in with_client
    assert "DL-1/General/M1 Design/1.1 Spec/1.1_Spec-Area A-20260226-V1.pdf" in paths(session, project, lang="en")


def test_milestone_layout_and_scope(session, project, d, stored):
    p = paths(session, project, layout="milestone")
    assert sum("72&73_" in x for x in p) == 1  # 同一个文件夹里只放一份
    assert "DL-1/M1 Design/1.2 Datasheet/1.2.72_Datasheet-Mixer-Wiring-20260226-V1.dwg" in p
    only73 = paths(session, project, scope=delivery.Scope(equipment_ids=[d["eq"]["73"]]))
    assert len(only73) == 1 and only73[0].startswith("DL-1/Alpha/73 Dryer/1.2.72&73_")
    only_spec = paths(session, project, scope=delivery.Scope(document_type_ids=[d["docs"][0]]))
    assert len(only_spec) == 1


def test_long_paths_shortened(session, project, d, api_client):
    store(api_client, project, [d["spec"]], ("Very long title " * 12) + ".pdf", b"long")
    p = paths(session, project)
    assert p and all(len(x) <= delivery.MAX_PATH for x in p)


# ---------------------------------------------------------------- 交付包


def build(api_client, project, session, body):
    pkg = call(api_client, "POST", project, "packages", body, 201)
    assert pkg["status"] == "queued"
    assert worker.run_once(session, "w", [delivery.PACKAGE_KIND])
    pkg = next(x for x in call(api_client, "GET", project, "packages") if x["id"] == pkg["id"])
    assert pkg["status"] == "succeeded" and pkg["available"], pkg
    resp = call(api_client, "GET", project, f"packages/{pkg['id']}/download")
    return pkg, zipfile.ZipFile(io.BytesIO(resp.content))


def test_package_zip_and_manifest(api_client, project, d, stored, session):
    est = call(api_client, "POST", project, "packages/estimate", {"lang": "zh"})
    assert est["files"] == 6 and est["size"] > 0  # Spec、72 两份、共用的两处、Layout
    pkg, z = build(api_client, project, session, {"lang": "zh"})
    names = z.namelist()
    assert "DL-1/文件清单.xlsx" in names and len(names) == 7
    assert z.read("DL-1/Alpha/72 Mixer/1.2.72_Datasheet-Mixer-Wiring-20260226-V1.dwg") == b"wiring"
    wb = load_workbook(io.BytesIO(z.read("DL-1/文件清单.xlsx")))
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 6 and all(r[5] for r in rows)  # 每行都有大小
    assert pkg["download_name"].startswith("DL-1_package_") and pkg["download_name"].endswith(".zip")
    call(api_client, "POST", project, "packages", {"scope": {"equipment_ids": [999999]}}, 422)


def test_package_as_submission(api_client, project, d, stored, session):
    pkg, _ = build(api_client, project, session, {
        "scope": {"document_type_ids": [d["docs"][1]]}, "submit": {"subject": "Datasheets"}})
    assert pkg["transmittal_no"]
    t = next(x for x in call(api_client, "GET", project, "transmittals") if x["number"] == pkg["transmittal_no"])
    detail = call(api_client, "GET", project, f"transmittals/{t['id']}")
    assert sorted(i["deliverable_id"] for i in detail["items"]) == sorted([d["72"], d["73"]])
    item72 = session.scalar(select(TransmittalItem).where(TransmittalItem.transmittal_id == t["id"],
                                                          TransmittalItem.deliverable_id == d["72"]))
    recorded = set(session.scalars(select(TransmittalItemVersion.version_id)
                                   .where(TransmittalItemVersion.item_id == item72.id)))
    assert recorded == {stored["ds"]["version_id"], stored["wiring"]["version_id"], stored["shared"]["version_id"]}
    assert files_svc.protected_version_ids(session, list(recorded)) == recorded
    items = {i["id"]: i for i in call(api_client, "GET", project, "deliverables")["items"]}
    assert items[d["72"]]["status"] == "submitted"


def test_transmittal_records_all_current_versions(api_client, project, d, stored, session):
    t = call(api_client, "POST", project, "transmittals", {"deliverable_ids": [d["72"]]}, 201)
    item = session.scalar(select(TransmittalItem).where(TransmittalItem.transmittal_id == t["id"]))
    recorded = set(session.scalars(select(TransmittalItemVersion.version_id)
                                   .where(TransmittalItemVersion.item_id == item.id)))
    assert len(recorded) == 3  # Datasheet、Wiring、共用的 Shared
    # 批准后其中一份（Wiring）又传了新版本：闭环检查挡关闭
    call(api_client, "POST", project, f"transmittals/{t['id']}/reply", {"code": "1", "reply_date": "2026-10-20"})
    issues = lambda: next(i for i in call(api_client, "GET", project, f"milestones/{d['m1']}/closeout")["items"]  # noqa: E731
                          if i["id"] == d["72"])["issues"]
    assert "new_version_after_approval" not in issues()
    store(api_client, project, [d["72"]], "Wiring V2 20260301.dwg", b"wiring 2")
    assert "new_version_after_approval" in issues()


# ---------------------------------------------------------------- 基线


def test_baseline_freeze_compare_export(api_client, project, d, stored, session):
    b = call(api_client, "POST", project, "baselines", {"name": "Design freeze", "description": "before review"}, 201)
    assert b["counts"]["deliverables"] == 4 and b["counts"]["missing"] == 0  # Spec、72、73（有共用的那份）、Layout A
    detail = call(api_client, "GET", project, f"baselines/{b['id']}")
    row72 = next(r for r in detail["items"] if r["deliverable_id"] == d["72"])
    assert len(row72["files"]) == 3 and row72["status"] == "not_submitted"
    # 之后：72 的 Datasheet 升版、Spec 提交
    store(api_client, project, [d["72"]], "Datasheet V2 20260301.pdf", b"ds 2")
    call(api_client, "POST", project, "transmittals", {"deliverable_ids": [d["spec"]]}, 201)
    diff = call(api_client, "GET", project, f"baselines/{b['id']}/compare")
    kinds = {r["deliverable_id"]: {c["kind"] for c in r["changes"]} for r in diff["rows"]}
    assert kinds[d["72"]] == {"new_version"} and kinds[d["spec"]] == {"status"}
    b2 = call(api_client, "POST", project, "baselines", {"name": "After review"}, 201)
    diff = call(api_client, "GET", project, f"baselines/{b['id']}/compare?other={b2['id']}")
    assert diff["counts"]["new_version"] == 1 and diff["counts"]["status"] >= 1
    assert call(api_client, "GET", project, f"baselines/{b2['id']}/compare")["rows"] == []  # 和现在一样
    wb = load_workbook(io.BytesIO(call(api_client, "GET", project, f"baselines/{b['id']}/export.xlsx").content))
    assert wb.active.cell(1, 1).value == "Design freeze"
    # 按基线导出交付包：放的是冻结时的版本
    _, z = build(api_client, project, session, {"baseline_id": b["id"], "lang": "zh"})
    assert z.read("DL-1/Alpha/72 Mixer/1.2.72_Datasheet-Mixer-Datasheet-20260226-V1.pdf") == b"ds"
    assert [x["name"] for x in call(api_client, "GET", project, "baselines")] == ["After review", "Design freeze"]


def test_baseline_protects_versions(api_client, project, d, stored, session):
    call(api_client, "POST", project, "baselines", {"name": "B1"}, 201)
    for n in range(2, 6):
        store(api_client, project, [d["spec"]], f"Area A V{n} 20260226.pdf", f"spec {n}".encode())
    v1 = session.get(Version, stored["spec"]["version_id"])
    assert v1.recycled_at is None  # 被基线引用：不进回收区


# ---------------------------------------------------------------- 镜像


@pytest.fixture
def mirror_dir(tmp_path, monkeypatch):
    from app.config import get_settings

    root = tmp_path / "mirror"
    monkeypatch.setattr(get_settings(), "mirror_dir", root)
    monkeypatch.setattr(get_settings(), "mirror_mode", "link")
    return root


def test_mirror_sync_links_renames_and_keeps_foreign_files(api_client, project, d, stored, session, mirror_dir):
    from app.models import Project

    p = session.get(Project, project["id"])
    counts = mirror.sync_project(session, p)
    assert counts["added"] == 7  # 6 处（共用的那份放两处）+ 客户文件
    f = mirror_dir / "DL-1/Alpha/72 Mixer/1.2.72_Datasheet-Mixer-Wiring-20260226-V1.dwg"
    assert f.read_bytes() == b"wiring"
    assert not os.access(f, os.W_OK)  # 只读
    if counts["linked"]:
        assert os.stat(f).st_nlink >= 2  # 同一块盘：硬链接，不占空间
    (mirror_dir / "DL-1" / "my notes.txt").write_text("not ours")
    # 设备改名：旧路径删掉、新路径出现；别人放的文件不动
    call(api_client, "PATCH", project, f"equipment/{d['eq']['72']}", {"name": "Mixer A"})
    counts = mirror.sync_project(session, p)
    assert not f.exists()
    assert (mirror_dir / "DL-1/Alpha/72 Mixer A/1.2.72_Datasheet-Mixer A-Wiring-20260226-V1.dwg").is_file()
    assert (mirror_dir / "DL-1" / "my notes.txt").is_file()
    assert not (mirror_dir / "DL-1/Alpha/72 Mixer").exists()  # 空文件夹删掉
    assert counts["kept"] >= 1
    # 删除平台里的文件实体时，只读的硬链接不挡路
    sha = session.get(Version, stored["layout"]["version_id"]).blob_sha256
    get_storage().delete(sha)
    assert not get_storage().exists(sha)


def test_mirror_copy_mode(api_client, project, d, stored, session, mirror_dir, monkeypatch):
    from app.config import get_settings
    from app.models import Project

    monkeypatch.setattr(get_settings(), "mirror_mode", "copy")
    counts = mirror.sync_project(session, session.get(Project, project["id"]))
    assert counts["copied"] == counts["added"] and counts["linked"] == 0
    f = mirror_dir / "DL-1/通用文件/M1 Design/1.1 Spec/1.1_Spec-Area A-20260226-V1.pdf"
    assert os.stat(f).st_nlink == 1 and not (os.stat(f).st_mode & stat.S_IWRITE)


def test_mirror_job_follows_events(api_client, project, d, stored, session, mirror_dir):
    job = mirror.ensure_job(session)
    assert job is not None and mirror.ensure_job(session) is None  # 已经排着一个
    assert worker.run_once(session, "w", [mirror.MIRROR_KIND])
    session.refresh(job)
    assert job.result["projects"] >= 1 and mirror.pending(session) == 0
    assert mirror.ensure_job(session) is None  # 没有新变化
    store(api_client, project, [d["73"]], "Dryer only 20260226.pdf", b"dryer")
    assert mirror.ensure_job(session) is not None
    worker.run_once(session, "w", [mirror.MIRROR_KIND])
    assert (mirror_dir / "DL-1/Alpha/73 Dryer").is_dir()
    st = call(api_client, "GET", project, "mirror")
    assert st["enabled"] and st["pending_events"] == 0 and st["last_job"]["status"] == "succeeded"
    call(api_client, "POST", project, "mirror/sync")
    assert session.scalar(select(Job).where(Job.kind == mirror.MIRROR_KIND, Job.status == "queued")) is not None


def test_mirror_disabled(api_client, project, session):
    assert mirror.ensure_job(session) is None or mirror.mirror_root() is not None
    st = call(api_client, "GET", project, "mirror")
    assert st["enabled"] is (mirror.mirror_root() is not None)
