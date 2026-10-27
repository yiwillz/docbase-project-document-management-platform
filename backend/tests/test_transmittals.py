"""第 8B 步：审核代码、传送单（生成 / 补记 / 作废）、客户回复、导出与导回、MDL 导回补记录。

虚构项目 TX-1（客户审核期限 14 天）：
M1（截止 2099-01-01）
  1.1 整个种类 Spec（合同规定，依据 App D 3.1）
  1.2 按模块 A、B Layout（合同规定，没写依据）
  1.3 整个种类 Noise study（客户要求，依据 CL-015，待确认）
  1.4 整个种类 Checklist（内部要求）
"""

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import Deliverable, ReviewCode, Transmittal, TransmittalItem
from app.services import mdl_export as mx
from app.services import review_codes
from app.services import transmittal_export as tx

USER = {"X-User": "alice"}


def _url(project, path=""):
    return f"/api/projects/{project['id']}/{path}"


def call(api_client, method, project, path, body=None, status=200, **kwargs):
    resp = api_client.request(method, _url(project, path), json=body, headers=USER, **kwargs)
    assert resp.status_code == status, resp.text
    return resp.json() if resp.headers.get("content-type", "").startswith("application/json") else resp


@pytest.fixture
def project(api_client):
    return api_client.post("/api/projects", json={
        "code": "TX-1", "name": "Transmittal demo", "client_name": "Client Co.",
        "contractor_name": "Contractor Co.", "review_days": 14}, headers=USER).json()


@pytest.fixture
def demo(api_client, project, session):
    m1 = call(api_client, "POST", project, "milestones", {"name": "Design", "due_date": "2099-01-01"}, 201)
    docs = call(api_client, "POST", project, "document-types/bulk", {
        "milestone_id": m1["id"], "rows": [
            {"name": "Spec", "basis": "App D 3.1"}, {"name": "Layout"},
            {"name": "Noise study", "origin": "client", "basis": "CL-015",
             "required_date": "2026-09-01"},
            {"name": "Checklist", "origin": "internal"}]}, 201)
    call(api_client, "PUT", project, f"document-types/{docs[1]['id']}/modules", {"module_names": ["A", "B"]})
    modules = {m["name"]: m["id"] for m in call(api_client, "GET", project, "modules")}
    keys = {}
    for d in session.scalars(select(Deliverable).where(Deliverable.project_id == project["id"],
                                                       Deliverable.obsolete.is_(False))):
        no = next(x["submittal_no"] for x in docs if x["id"] == d.document_type_id)
        if d.module_id:
            no += "/" + next(n for n, i in modules.items() if i == d.module_id)
        keys[no] = d.id
    return {"m1": m1, "docs": docs, "d": keys}


def items(api_client, project, **params):
    return {it["id"]: it for it in call(api_client, "GET", project, "deliverables",
                                        params=params)["items"]}


def submit(api_client, project, ids, **body):
    return call(api_client, "POST", project, "transmittals", {"deliverable_ids": ids, **body}, 201)


# ---------------------------------------------------------------- 审核代码


def test_new_project_gets_default_review_codes(api_client, project):
    codes = call(api_client, "GET", project, "review-codes")
    assert [(c["code"], c["status"], c["final"]) for c in codes] == [
        ("1", "approved", True), ("2", "approved", False), ("3", "revise_resubmit", False),
        ("4", "revise_resubmit", False), ("5", "approved", True)]
    assert codes[1]["label_zh"] == "带意见批准"


def test_replace_review_codes_and_find(api_client, project, session):
    body = {"codes": [
        {"code": "A", "label_en": "Approved", "status": "approved"},
        {"code": "B", "label_en": "Approved as noted", "status": "approved", "final": False,
         "aliases": ["AAN"]},
        {"code": "C", "label_en": "Revise & Resubmit", "status": "revise_resubmit", "final": False}],
        "reason": "客户程序"}
    codes = call(api_client, "PUT", project, "review-codes", body)
    assert [c["code"] for c in codes] == ["A", "B", "C"]
    found = review_codes.list_codes(session, project["id"])
    assert review_codes.find_code(found, "aan").code == "B"
    assert review_codes.find_code(found, "B - Approved as noted").code == "B"
    assert review_codes.find_code(found, "approved").code == "A"
    assert review_codes.find_code(found, "X") is None
    old = session.scalars(select(ReviewCode).where(ReviewCode.project_id == project["id"],
                                                   ReviewCode.active.is_(False))).all()
    assert len(old) == 5  # 旧的停用，留历史
    dup = {"codes": [{"code": "A", "label_en": "x", "status": "approved"},
                     {"code": "a", "label_en": "y", "status": "approved"}]}
    assert api_client.put(_url(project, "review-codes"), json=dup, headers=USER).status_code == 422
    bad = {"codes": [{"code": "A", "label_en": "x", "status": "not_submitted"}]}
    assert api_client.put(_url(project, "review-codes"), json=bad, headers=USER).status_code == 422


# ---------------------------------------------------------------- 生成、补记


def test_create_transmittal_marks_submitted(api_client, project, demo):
    d = demo["d"]
    t = submit(api_client, project, [d["1.1"], d["1.2/A"]], subject="Design docs",
               submitted_date="2026-10-20", revisions={str(d["1.1"]): "Rev A"})
    assert (t["number"], t["source"], t["review_due"]) == ("TX-1-TR-0001", "platform", "2026-11-03")
    assert [i["revision"] for i in t["items"]] == ["Rev A", None]
    rows = items(api_client, project)
    for key in ("1.1", "1.2/A"):
        it = rows[d[key]]
        assert (it["status"], it["responsible"], it["submitted_date"], it["status_date"]) == (
            "submitted", "client", "2026-10-20", "2026-10-20")
        assert (it["transmittal_no"], it["status_basis"]) == ("TX-1-TR-0001", "传送单 TX-1-TR-0001")
    assert rows[d["1.2/B"]]["status"] == "not_submitted"
    # 编号接着排；前缀可在项目编号规则里改
    call(api_client, "PATCH", project, "", {"numbering_rules": {"transmittal_prefix": "T"}})
    assert submit(api_client, project, [d["1.3"]])["number"] == "TX-1-T-0002"
    listed = call(api_client, "GET", project, "transmittals")
    assert [(x["number"], x["items"], x["replied"]) for x in listed] == [
        ("TX-1-T-0002", 1, 0), ("TX-1-TR-0001", 2, 0)]


def test_platform_transmittal_defaults_to_today_record_needs_date(api_client, project, demo):
    t = submit(api_client, project, [demo["d"]["1.1"]])
    assert t["submitted_date"] == date.today().isoformat()
    resp = api_client.post(_url(project, "transmittals/record"), headers=USER,
                           json={"deliverable_ids": [demo["d"]["1.2/A"]]})
    assert resp.status_code == 422 and resp.json()["detail"]["code"] == "transmittal.date_required"
    t = call(api_client, "POST", project, "transmittals/record", {
        "deliverable_ids": [demo["d"]["1.2/A"]], "submitted_date": "2026-09-01",
        "method": "email", "external_ref": "MAIL-77"}, 201)
    assert (t["source"], t["method"], t["external_ref"]) == ("recorded", "email", "MAIL-77")


def test_create_rejects_bad_deliverables(api_client, project, demo, session):
    old = session.scalars(select(Deliverable).where(Deliverable.project_id == project["id"],
                                                    Deliverable.obsolete.is_(True))).first()
    resp = api_client.post(_url(project, "transmittals"), headers=USER,
                           json={"deliverable_ids": [old.id]})
    assert resp.status_code == 409
    resp = api_client.post(_url(project, "transmittals"), headers=USER,
                           json={"deliverable_ids": [demo["d"]["1.1"], 999999]})
    assert resp.status_code == 404
    assert session.scalars(select(Transmittal).where(Transmittal.project_id == project["id"])).all() == []


# ---------------------------------------------------------------- 客户回复


def test_reply_whole_transmittal_with_override(api_client, project, demo):
    d = demo["d"]
    t = submit(api_client, project, [d["1.1"], d["1.2/A"], d["1.2/B"]], submitted_date="2026-10-20")
    a_item = next(i["item_id"] for i in t["items"] if i["deliverable_id"] == d["1.2/A"])
    result = call(api_client, "POST", project, f"transmittals/{t['id']}/reply", {
        "code": "1", "reply_date": "2026-11-02", "reply_ref": "CL-021",
        "items": [{"item_id": a_item, "code": "3", "reply_comments": "Fix dims"}]})
    assert result == {"updated": 3, "unchanged": 0}
    rows = items(api_client, project)
    spec, a = rows[d["1.1"]], rows[d["1.2/A"]]
    assert (spec["status"], spec["responsible"], spec["status_date"], spec["review_code"]) == (
        "approved", "none", "2026-11-02", "1")
    assert spec["status_basis"] == "审核代码 1，CL-021（传送单 TX-1-TR-0001）"
    assert (spec["review_final"], spec["no_evidence"], spec["reply_ref"]) == (True, False, "CL-021")
    assert (a["status"], a["responsible"], a["review_code"]) == ("revise_resubmit", "ours", "3")
    detail = call(api_client, "GET", project, f"transmittals/{t['id']}")
    a_row = next(i for i in detail["items"] if i["item_id"] == a_item)
    assert (a_row["reply_comments"], a_row["reply_source"], a_row["replied_by"]) == (
        "Fix dims", "manual", "alice")
    # 重交 → 已提交；再回复带意见批准 → 已批准但不是最终
    second = submit(api_client, project, [d["1.2/A"]], submitted_date="2026-11-10")
    assert items(api_client, project)[d["1.2/A"]]["status"] == "submitted"
    call(api_client, "POST", project, f"transmittals/{second['id']}/reply",
         {"code": "Approved with comments", "reply_date": "2026-11-15"})
    a = items(api_client, project)[d["1.2/A"]]
    assert (a["status"], a["review_code"], a["review_final"]) == ("approved", "2", False)


def test_reply_validation_and_deemed(api_client, project, demo):
    t = submit(api_client, project, [demo["d"]["1.1"]], submitted_date="2026-10-01")
    url = f"transmittals/{t['id']}/reply"
    resp = api_client.post(_url(project, url), headers=USER, json={"code": "9"})
    assert resp.status_code == 422 and resp.json()["detail"]["code"] == "transmittal.bad_code"
    resp = api_client.post(_url(project, url), headers=USER, json={"code": "1", "deemed": True})
    assert resp.json()["detail"]["code"] == "transmittal.deemed_needs_basis"
    resp = api_client.post(_url(project, url), headers=USER, json={})
    assert resp.json()["detail"]["code"] == "transmittal.no_reply"
    call(api_client, "POST", project, url, {"code": "1", "deemed": True,
                                            "deemed_basis": "Contract 8.4 (14 days)"})
    it = items(api_client, project)[demo["d"]["1.1"]]
    assert it["status"] == "approved" and "视为批准：Contract 8.4 (14 days)" in it["status_basis"]
    # 撤回回复：回到已提交
    call(api_client, "POST", project, url, {"code": ""})
    assert items(api_client, project)[demo["d"]["1.1"]]["status"] == "submitted"


def test_void_recomputes_from_remaining_records(api_client, project, demo):
    d = demo["d"]["1.1"]
    first = submit(api_client, project, [d], submitted_date="2026-10-01")
    call(api_client, "POST", project, f"transmittals/{first['id']}/reply",
         {"code": "3", "reply_date": "2026-10-05"})
    second = submit(api_client, project, [d], submitted_date="2026-10-10")
    assert items(api_client, project)[d]["status"] == "submitted"
    resp = api_client.post(_url(project, f"transmittals/{second['id']}/void"), headers=USER, json={"reason": " "})
    assert resp.status_code == 422
    call(api_client, "POST", project, f"transmittals/{second['id']}/void", {"reason": "发错了"})
    it = items(api_client, project)[d]
    assert (it["status"], it["transmittal_no"]) == ("revise_resubmit", first["number"])
    call(api_client, "POST", project, f"transmittals/{first['id']}/void", {"reason": "重复"})
    it = items(api_client, project)[d]
    assert (it["status"], it["submitted_date"], it["transmittal_no"]) == ("not_submitted", None, None)
    assert [x["number"] for x in call(api_client, "GET", project, "transmittals")] == []
    assert len(call(api_client, "GET", project, "transmittals", params={"include_void": True})) == 2
    resp = api_client.post(_url(project, f"transmittals/{first['id']}/reply"), headers=USER, json={"code": "1"})
    assert resp.status_code == 409


def test_header_date_change_resyncs(api_client, project, demo):
    t = submit(api_client, project, [demo["d"]["1.1"]], submitted_date="2026-10-01")
    call(api_client, "PATCH", project, f"transmittals/{t['id']}", {"submitted_date": "2026-09-28"})
    it = items(api_client, project)[demo["d"]["1.1"]]
    assert (it["submitted_date"], it["review_due"]) == ("2026-09-28", "2026-10-12")


def test_review_overdue_and_no_evidence(api_client, project, demo):
    d = demo["d"]
    submit(api_client, project, [d["1.1"]], submitted_date="2020-01-01")  # 早就超过 14 天
    call(api_client, "PATCH", project, f"deliverables/{d['1.2/A']}", {"status": "approved"})
    call(api_client, "PATCH", project, f"deliverables/{d['1.2/B']}", {"status": "not_applicable"})
    call(api_client, "PATCH", project, f"deliverables/{d['1.3']}",
         {"status": "not_applicable", "status_basis": "客户函 CL-030 取消"})
    rows = items(api_client, project)
    assert rows[d["1.1"]]["review_overdue"] is True
    assert [rows[d[k]]["no_evidence"] for k in ("1.2/A", "1.2/B", "1.3")] == [True, True, False]
    s = call(api_client, "GET", project, "deliverables/summary")
    assert (s["overall"]["review_overdue"], s["overall"]["no_evidence"]) == (1, 2)
    assert list(items(api_client, project, review_overdue=True)) == [d["1.1"]]
    assert sorted(items(api_client, project, no_evidence=True)) == sorted([d["1.2/A"], d["1.2/B"]])


# ---------------------------------------------------------------- 导出、客户填好导回


def _export(api_client, project, t, **params):
    return call(api_client, "GET", project, f"transmittals/{t['id']}/export", params=params).content


def test_export_excel_and_pdf(api_client, project, demo):
    d = demo["d"]
    t = submit(api_client, project, [d["1.2/B"], d["1.1"], d["1.2/A"]], subject="For review",
               submitted_date="2026-10-20")
    wb = load_workbook(io.BytesIO(_export(api_client, project, t)))
    ws = wb["Transmittal"]
    assert ws["A1"].value == "DOCUMENT TRANSMITTAL" and ws["G2"].value == "TX-1-TR-0001"
    assert (ws["B3"].value, ws["B5"].value) == ("Client Co.", "For review")
    rows = [[ws.cell(r, c).value for c in (2, 3, 4)] for r in range(8, 11)]
    assert rows == [["1.1", "Spec", None], ["1.2", "Layout", "A"], ["1.2", "Layout", "B"]]
    assert ws.cell(8, tx.C_CODE).fill.fgColor.rgb.endswith(tx.FILL_IN)
    assert wb["_meta"].sheet_state == "hidden" and ws.column_dimensions["K"].hidden
    zh = load_workbook(io.BytesIO(_export(api_client, project, t, lang="zh")))
    assert zh.worksheets[0]["A1"].value == "文件传送单"
    for lang in ("en", "zh"):
        pdf = _export(api_client, project, t, format="pdf", lang=lang)
        assert pdf.startswith(b"%PDF") and b"docbase \\(system generated\\)" in pdf


def _fill_and_preview(api_client, project, t, fill, name="reply.xlsx"):
    wb = load_workbook(io.BytesIO(_export(api_client, project, t)))
    meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(max_col=2)}
    fill(wb.worksheets[0], meta)
    buffer = io.BytesIO()
    wb.save(buffer)
    resp = api_client.post(_url(project, "transmittals/import/preview"), headers=USER,
                           files={"file": (name, buffer.getvalue())})
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_transmittal_return_round_trip(api_client, project, demo, session):
    d = demo["d"]
    t = submit(api_client, project, [d["1.1"], d["1.2/A"], d["1.2/B"]], submitted_date="2026-10-20")

    def fill(ws, meta):
        ws.cell(8, tx.C_CODE, "1")
        ws.cell(8, tx.C_DATE, date(2026, 11, 3))
        ws.cell(9, tx.C_CODE, "Revise & Resubmit")
        ws.cell(9, tx.C_DATE, "2026-11-03")
        ws.cell(9, tx.C_COMMENTS, "See markup")
        ws[meta["ref_cell"]] = "CT-0042"

    preview = _fill_and_preview(api_client, project, t, fill)
    assert preview["ok"] and preview["reply_ref"] == "CT-0042" and len(preview["changes"]) == 2
    assert preview["warnings"] == []
    counts = call(api_client, "POST", project, "transmittals/import/confirm", {"token": preview["token"]})
    assert counts == {"updated": 2, "unchanged": 0}
    rows = items(api_client, project)
    assert (rows[d["1.1"]]["status"], rows[d["1.1"]]["reply_ref"]) == ("approved", "CT-0042")
    assert rows[d["1.2/A"]]["status"] == "revise_resubmit"
    assert rows[d["1.2/B"]]["status"] == "submitted"  # 没填的不动
    item = session.scalars(select(TransmittalItem).where(TransmittalItem.deliverable_id == d["1.2/A"])).one()
    assert (item.reply_source.value, item.reply_comments) == ("transmittal_return", "See markup")
    # 原样再导回：没有差异；没填回函号时用文件名
    again = _fill_and_preview(api_client, project, t, lambda ws, meta: None, name="again.xlsx")
    assert again["changes"] == [] and again["reply_ref"] == "传送单导回（again.xlsx）"


def test_transmittal_return_errors(api_client, project, demo):
    t = submit(api_client, project, [demo["d"]["1.1"]], submitted_date="2026-10-20")
    preview = _fill_and_preview(api_client, project, t, lambda ws, meta: ws.cell(8, tx.C_CODE, "Z"))
    assert not preview["ok"] and preview["errors"][0]["code"] == "transmittal_import.bad_code"
    resp = api_client.post(_url(project, "transmittals/import/confirm"), headers=USER,
                           json={"token": preview["token"]})
    assert resp.json()["detail"]["code"] == "transmittal_import.has_errors"
    no_date = _fill_and_preview(api_client, project, t, lambda ws, meta: ws.cell(8, tx.C_CODE, "1"))
    assert no_date["ok"] and no_date["warnings"][0]["code"] == "transmittal_import.no_reply_date"
    other = api_client.post("/api/projects", json={"code": "TX-2", "name": "o"}, headers=USER).json()
    data = _export(api_client, project, t)
    resp = api_client.post(_url(other, "transmittals/import/preview"), headers=USER,
                           files={"file": ("x.xlsx", data)}).json()
    assert resp["errors"][0]["code"] == "transmittal_import.other_project"
    buffer = io.BytesIO()
    Workbook().save(buffer)
    resp = api_client.post(_url(project, "transmittals/import/preview"), headers=USER,
                           files={"file": ("y.xlsx", buffer.getvalue())}).json()
    assert resp["errors"][0]["code"] == "transmittal_import.not_platform_file"
    call(api_client, "POST", project, f"transmittals/{t['id']}/void", {"reason": "x"})
    resp = api_client.post(_url(project, "transmittals/import/preview"), headers=USER,
                           files={"file": ("z.xlsx", data)}).json()
    assert resp["errors"][0]["code"] == "transmittal_import.void"


# ---------------------------------------------------------------- MDL 导回补记录


def test_mdl_return_creates_reply_records(api_client, project, demo, session):
    d = demo["d"]
    t = submit(api_client, project, [d["1.1"]], submitted_date="2026-10-01")
    wb = load_workbook(io.BytesIO(call(api_client, "GET", project, "mdl/export").content))
    ws = wb["MDL"]
    meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(max_col=2)}
    rows = range(int(meta["header_row"]) + 1, int(meta["last_row"]) + 1)
    r11 = next(r for r in rows if ws.cell(r, mx.C_NO).value == "1.1")
    r13 = next(r for r in rows if ws.cell(r, mx.C_NO).value == "1.3")
    ws.cell(r11, mx.C_STATUS, "Approved")
    ws.cell(r11, mx.C_SDATE, date(2026, 10, 9))
    ws.cell(r13, mx.C_STATUS, "Revise & Resubmit")
    buffer = io.BytesIO()
    wb.save(buffer)
    preview = api_client.post(_url(project, "mdl/import/preview"), headers=USER,
                              files={"file": ("client.xlsx", buffer.getvalue())}).json()
    call(api_client, "POST", project, "mdl/import/confirm", {"token": preview["token"]})
    session.expire_all()
    item = session.scalars(select(TransmittalItem).where(TransmittalItem.transmittal_id == t["id"])).one()
    assert (item.review_code, item.reply_date, item.reply_ref, item.reply_source.value) == (
        "1", date(2026, 10, 9), "MDL 导回（client.xlsx）", "mdl_return")
    auto = session.scalars(select(Transmittal).where(Transmittal.project_id == project["id"],
                                                     Transmittal.source == "mdl_return")).one()
    auto_item = session.scalars(select(TransmittalItem).where(TransmittalItem.transmittal_id == auto.id)).one()
    assert (auto_item.deliverable_id, auto_item.review_code) == (d["1.3"], "3")
    rows = items(api_client, project)
    assert rows[d["1.1"]]["no_evidence"] is False and rows[d["1.1"]]["status"] == "approved"
