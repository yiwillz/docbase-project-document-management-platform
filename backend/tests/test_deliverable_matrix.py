"""第 8A 步：应交项列表、分组列表、矩阵、汇总、修改、缺件清单。

虚构项目（截止日用 2020 年和 2099 年，结果不随运行日期变化）：

M1（截止 2020-01-01，已过）
  1.1 整个种类                      已批准
  1.2 按模块 A、B                   A 未提交（逾期）、B 不适用
  1.3 按设备（分区 A、B，全部设备） a1 已批准、a2 需修改重交（我方）、b1 审核中（客户）、
                                    c1 未提交（逾期，责任方改为供应商）
  1.4 内部要求，整个种类            已提交（责任方未指定）
M2（截止 2099-01-01）
  2.1 整个种类，种类截止 2020-06-01 未提交（逾期，截止日来自种类）
  2.2 整个种类，应交项截止 2020-02-01 未提交（逾期，截止日来自应交项）
  2.3 整个种类                      未提交（不逾期）

设备：a1、a2 → 种类 TA（对应模块 A，供应商 V1）；b1 → TB（模块 B，供应商 V2）；
c1 → TC（没有对应模块，没有供应商）。
"""

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    ChangeLog,
    Deliverable,
    Equipment,
    EquipmentType,
    Event,
    ModuleEquipmentType,
)
from app.services import mdl_export as mx

USER = {"X-User": "alice"}


@pytest.fixture
def project(api_client):
    return api_client.post("/api/projects", json={"code": "DM-1", "name": "Matrix demo"},
                           headers=USER).json()


def _url(project, path=""):
    return f"/api/projects/{project['id']}/{path}"


def _call(api_client, method, project, path, body=None, status=200, **kwargs):
    resp = api_client.request(method, _url(project, path), json=body, headers=USER, **kwargs)
    assert resp.status_code == status, resp.text
    return resp.json() if resp.headers.get("content-type", "").startswith("application/json") else resp


@pytest.fixture
def demo(api_client, project, session):
    pid = project["id"]
    m1 = _call(api_client, "POST", project, "milestones", {"name": "Design", "due_date": "2020-01-01"}, 201)
    m2 = _call(api_client, "POST", project, "milestones", {"name": "FAT", "due_date": "2099-01-01"}, 201)
    d11, d12, d13, d14 = _call(api_client, "POST", project, "document-types/bulk", {
        "milestone_id": m1["id"], "rows": [{"name": "Spec"}, {"name": "Layout"},
                                           {"name": "Datasheet"},
                                           {"name": "Checklist", "origin": "internal"}]}, 201)
    d21, d22, d23 = _call(api_client, "POST", project, "document-types/bulk", {
        "milestone_id": m2["id"], "rows": [{"name": "FAT procedure", "due_date": "2020-06-01"},
                                           {"name": "FAT report"}, {"name": "Punch list"}]}, 201)
    for doc in (d12, d13):
        _call(api_client, "PUT", project, f"document-types/{doc['id']}/modules", {"module_names": ["A", "B"]})
    modules = {m["name"]: m["id"] for m in _call(api_client, "GET", project, "modules")}
    ta, tb, tc = (EquipmentType(project_id=pid, name=n) for n in ("TA", "TB", "TC"))
    session.add_all([ta, tb, tc])
    session.flush()
    eqs = {tag: Equipment(project_id=pid, equipment_type_id=t.id, tag=tag, name=f"Tool {tag}",
                          vendor=v, sort_order=i)
           for i, (tag, t, v) in enumerate((("a1", ta, "V1"), ("a2", ta, "V1"), ("b1", tb, "V2"),
                                            ("c1", tc, None)))}
    session.add_all(eqs.values())
    session.add_all([ModuleEquipmentType(module_id=modules["A"], equipment_type_id=ta.id),
                     ModuleEquipmentType(module_id=modules["B"], equipment_type_id=tb.id)])
    session.commit()
    _call(api_client, "POST", project, "document-types/bulk-equipment",
          {"document_type_ids": [d13["id"]], "all_equipment": True, "mode": "add"})

    rows = {}
    for d in session.scalars(select(Deliverable).where(Deliverable.project_id == pid,
                                                       Deliverable.obsolete.is_(False))):
        doc = {d11["id"]: "1.1", d12["id"]: "1.2", d13["id"]: "1.3", d14["id"]: "1.4",
               d21["id"]: "2.1", d22["id"]: "2.2", d23["id"]: "2.3"}[d.document_type_id]
        if d.equipment_id:
            doc += "/" + next(t for t, e in eqs.items() if e.id == d.equipment_id)
        elif d.module_id:
            doc += "/" + next(n for n, i in modules.items() if i == d.module_id)
        rows[doc] = d.id

    def bulk(keys, changes):
        _call(api_client, "POST", project, "deliverables/bulk-update",
              {"ids": [rows[k] for k in keys], "changes": changes, "reason": "setup"})

    bulk(["1.1", "1.3/a1"], {"status": "approved"})
    bulk(["1.2/B"], {"status": "not_applicable"})
    bulk(["1.3/a2"], {"status": "revise_resubmit"})
    bulk(["1.3/b1"], {"status": "under_review"})
    bulk(["1.3/c1"], {"responsible": "vendor"})
    bulk(["1.4"], {"status": "submitted"})
    bulk(["2.2"], {"due_date": "2020-02-01"})
    return {"rows": rows, "modules": modules, "m1": m1, "m2": m2, "eqs": eqs}


def _by_key(items, demo):
    names = {v: k for k, v in demo["rows"].items()}
    return {names[it["id"]]: it for it in items}


# ---------------------------------------------------------------- 列表


def test_list_returns_every_deliverable_with_derived_fields(api_client, project, demo):
    page = _call(api_client, "GET", project, "deliverables")
    assert page["total"] == 11
    items = _by_key(page["items"], demo)
    assert list(items) == ["1.1", "1.2/A", "1.2/B", "1.3/a1", "1.3/a2", "1.3/b1", "1.3/c1",
                           "1.4", "2.1", "2.2", "2.3"]  # 里程碑 → 种类 → 模块 → 设备
    a1, c1 = items["1.3/a1"], items["1.3/c1"]
    assert (a1["module_name"], a1["equipment_type_name"], a1["vendor"]) == ("A", "TA", "V1")
    assert (c1["module_name"], c1["vendor"], c1["overdue"]) == (None, None, True)
    assert [(items[k]["due_date"], items[k]["due_source"], items[k]["overdue"]) for k in
            ("1.1", "2.1", "2.2", "2.3")] == [
        ("2020-01-01", "milestone", False),  # 已批准不算逾期
        ("2020-06-01", "document_type", True),
        ("2020-02-01", "deliverable", True),
        ("2099-01-01", "milestone", False)]
    assert items["1.3/a2"]["responsible"] == "ours"  # 改状态时带出
    assert items["1.4"]["origin"] == "internal" and not items["1.4"]["in_client_mdl"]


@pytest.mark.parametrize(("query", "expected"), [
    ("status=not_submitted&status=revise_resubmit",
     ["1.2/A", "1.3/a2", "1.3/c1", "2.1", "2.2", "2.3"]),
    ("overdue=true", ["1.2/A", "1.3/c1", "2.1", "2.2"]),
    ("module_id={A}", ["1.2/A", "1.3/a1", "1.3/a2"]),
    ("vendor=V1", ["1.3/a1", "1.3/a2"]),
    ("vendor=", ["1.1", "1.2/A", "1.2/B", "1.3/c1", "1.4", "2.1", "2.2", "2.3"]),
    ("responsible=client", ["1.3/b1"]),
    ("origin=internal", ["1.4"]),
    ("in_client_mdl=false", ["1.4"]),
    ("kind=per_equipment&milestone_id={m1}", ["1.3/a1", "1.3/a2", "1.3/b1", "1.3/c1"]),
    ("q=tool b1", ["1.3/b1"]),
    ("q=fat", ["2.1", "2.2"]),
])
def test_list_filters(api_client, project, demo, query, expected):
    query = query.format(A=demo["modules"]["A"], m1=demo["m1"]["id"])
    page = _call(api_client, "GET", project, f"deliverables?{query}")
    assert list(_by_key(page["items"], demo)) == expected


def test_list_pagination(api_client, project, demo):
    page = _call(api_client, "GET", project, "deliverables?offset=8&limit=5")
    assert (page["total"], len(page["items"])) == (11, 3)


# ---------------------------------------------------------------- 汇总（和人工数的一致）


def test_summary_matches_hand_count(api_client, project, demo):
    s = _call(api_client, "GET", project, "deliverables/summary")
    # 作废 4 个：1.2、1.3 原来"整个种类"的，1.3 改按设备后原来按模块的 A、B
    assert (s["active"], s["obsolete"]) == (11, 4)
    assert s["by_kind"] == {"whole": 5, "per_module": 2, "per_equipment": 4}
    m1, m2 = s["milestones"]
    # M1：8 项，不适用 1，已批准 2 → 2/7；逾期 1.2/A、c1；待办：我方 a2、客户 b1、供应商 c1、未指定 1.2/A 和 1.4
    assert (m1["code"], m1["total"], m1["applicable"], m1["approved"], m1["overdue"]) == ("M1", 8, 7, 2, 2)
    assert m1["completion"] == round(2 / 7, 4)
    assert m1["pending"] == {"ours": 1, "vendor": 1, "client": 1, "unassigned": 2}
    assert (m2["total"], m2["approved"], m2["completion"], m2["overdue"]) == (3, 0, 0.0, 2)
    assert m2["pending"] == {"ours": 0, "vendor": 0, "client": 0, "unassigned": 3}
    modules = {m["module_name"]: m for m in s["modules"]}
    # 模块 A：1.2/A、a1、a2 → 1/3；模块 B：1.2/B（不适用）、b1 → 0/1；c1 和整个种类的没有模块
    assert (modules["A"]["total"], modules["A"]["approved"], modules["A"]["completion"]) == (3, 1, round(1 / 3, 4))
    assert (modules["B"]["total"], modules["B"]["applicable"], modules["B"]["completion"]) == (2, 1, 0.0)
    assert modules[None]["total"] == 6 and list(modules)[-1] is None
    m1_modules = {m["module_name"]: m["total"] for m in m1["modules"]}
    assert m1_modules == {"A": 3, "B": 2, None: 3}
    overall = s["overall"]
    assert (overall["total"], overall["approved"], overall["overdue"]) == (11, 2, 4)
    assert overall["completion"] == round(2 / 10, 4)


def test_summary_respects_filters(api_client, project, demo):
    s = _call(api_client, "GET", project, "deliverables/summary?in_client_mdl=true")
    assert s["overall"]["total"] == 10  # 不含内部要求


# ---------------------------------------------------------------- 分组列表、矩阵


def test_tree_groups_per_equipment_by_module(api_client, project, demo):
    t = _call(api_client, "GET", project, "deliverables/tree")
    m1 = t["milestones"][0]
    assert [d["submittal_no"] for d in m1["document_types"]] == ["1.1", "1.2", "1.3", "1.4"]
    d13 = m1["document_types"][2]
    assert [(g["module_name"], [i["equipment_tag"] for i in g["items"]]) for g in d13["modules"]] == [
        ("A", ["a1", "a2"]), ("B", ["b1"]), (None, ["c1"])]
    assert d13["stats"]["approved"] == 1 and d13["modules"][0]["stats"]["total"] == 2
    d12 = m1["document_types"][1]
    assert d12["modules"] == [] and [i["module_name"] for i in d12["items"]] == ["A", "B"]
    assert t["stats"]["total"] == 11


def test_matrix(api_client, project, demo):
    mtx = _call(api_client, "GET", project, f"deliverables/matrix?milestone_id={demo['m1']['id']}")
    assert [m["name"] for m in mtx["modules"]] == ["A", "B"]
    assert [(g["name"], [e["tag"] for e in g["equipment"]]) for g in mtx["equipment_groups"]] == [
        ("TA", ["a1", "a2"]), ("TB", ["b1"]), ("TC", ["c1"])]
    rows = {r["submittal_no"]: r for r in mtx["rows"]}
    assert list(rows) == ["1.1", "1.2", "1.3", "1.4"]
    assert list(rows["1.1"]["cells"]) == ["whole"]
    a = demo["modules"]["A"]
    assert rows["1.2"]["cells"][f"m:{a}"]["overdue"] is True
    c1 = rows["1.3"]["cells"][f"e:{demo['eqs']['c1'].id}"]
    assert (c1["status"], c1["responsible"], c1["current_version"]) == ("not_submitted", "vendor", None)
    assert len(rows["1.3"]["cells"]) == 4


# ---------------------------------------------------------------- 修改


def test_patch_one_records_history_and_event(api_client, project, demo, session):
    key = demo["rows"]["1.3/b1"]
    item = _call(api_client, "PATCH", project, f"deliverables/{key}", {
        "status": "approved", "status_date": "2026-10-20", "status_basis": "客户函 CL-012",
        "reason": "客户批准"})
    assert (item["status"], item["responsible"], item["status_date"], item["status_basis"]) == (
        "approved", "none", "2026-10-20", "客户函 CL-012")
    log = session.scalars(select(ChangeLog).where(ChangeLog.entity_type == "deliverables",
                                                  ChangeLog.entity_id == str(key))
                          .order_by(ChangeLog.id.desc())).first()
    assert log.reason == "客户批准" and log.changes["status"] == ["under_review", "approved"]
    event = session.scalars(select(Event).where(Event.event_type == "deliverable.status_changed")
                            .order_by(Event.id.desc())).first()
    assert event.payload["to"] == "approved" and event.payload["status_date"] == "2026-10-20"
    # 只改截止日：不动状态日期
    item = _call(api_client, "PATCH", project, f"deliverables/{key}", {"due_date": "2099-12-31"})
    assert (item["due_source"], item["status_date"]) == ("deliverable", "2026-10-20")
    item = _call(api_client, "PATCH", project, f"deliverables/{key}", {"due_date": None})
    assert item["due_source"] == "milestone"


def test_patch_rejects_bad_value_and_obsolete(api_client, project, demo, session):
    key = demo["rows"]["1.1"]
    resp = api_client.patch(_url(project, f"deliverables/{key}"), json={"status": "maybe"}, headers=USER)
    assert resp.status_code == 422
    old = session.scalars(select(Deliverable).where(Deliverable.project_id == project["id"],
                                                    Deliverable.obsolete.is_(True))).first()
    resp = api_client.patch(_url(project, f"deliverables/{old.id}"), json={"remarks": "x"}, headers=USER)
    assert resp.status_code == 409 and resp.json()["detail"]["code"] == "deliverable.obsolete"


def test_bulk_update_is_all_or_nothing(api_client, project, demo, session):
    other = api_client.post("/api/projects", json={"code": "DM-2", "name": "other"}, headers=USER).json()
    keys = [demo["rows"]["2.1"], demo["rows"]["2.3"]]
    resp = api_client.post(_url(project, "deliverables/bulk-update"), headers=USER, json={
        "ids": [*keys, 999999], "changes": {"status": "submitted"}})
    assert resp.status_code == 404
    session.expire_all()
    assert {session.get(Deliverable, k).status.value for k in keys} == {"not_submitted"}
    result = _call(api_client, "POST", project, "deliverables/bulk-update", {
        "ids": [*keys, keys[0]], "changes": {"status": "submitted", "submitted_date": "2026-10-22"},
        "reason": "传送单 TR-005"})
    assert result == {"updated": 2, "unchanged": 0}
    again = _call(api_client, "POST", project, "deliverables/bulk-update", {
        "ids": keys, "changes": {"status": "submitted"}})
    assert again == {"updated": 0, "unchanged": 2}
    assert other["id"] != project["id"]


def test_archived_project_is_read_only(api_client, project, demo):
    _call(api_client, "POST", project, "archive", {})
    resp = api_client.patch(_url(project, f"deliverables/{demo['rows']['1.1']}"),
                            json={"remarks": "x"}, headers=USER)
    assert resp.status_code == 409
    assert _call(api_client, "GET", project, "deliverables")["total"] == 11  # 仍可查看


# ---------------------------------------------------------------- 缺件清单


def _attach(api_client, project, demo, key, name):
    """给应交项存一个文件（收集情况变成"已收到"）。"""
    did = demo["rows"][key]
    r = api_client.post(f"/api/projects/{project['id']}/files/preview", data={"deliverable_ids": str(did)},
                        files={"file": (name, name.encode(), "application/pdf")})
    assert r.status_code == 200, r.text
    _call(api_client, "POST", project, f"files/uploads/{r.json()['token']}/confirm", {})


def test_outstanding_export_by_vendor(api_client, project, demo):
    # 缺件 = 没收到文件的（不含已批准、不适用）+ 需修改重交的：
    # b1（审核中）、1.4（已提交）收到了文件就不算；a2 需修改重交，收到了旧文件也还算
    for key in ("1.3/b1", "1.4", "1.3/a2"):
        _attach(api_client, project, demo, key, f"{key.replace('/', '-')}.pdf")
    resp = _call(api_client, "GET", project, "deliverables/outstanding.xlsx")
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Summary", "V1", "(No vendor)"]  # V2 只有收到了文件的 b1
    summary = [[c.value for c in row] for row in wb["Summary"].iter_rows(min_row=6, max_col=3)]
    assert summary == [["V1", 1, 0], ["(No vendor)", 5, 4], ["Total", 6, 4]]
    v1 = wb["V1"]
    assert [v1.cell(6, c).value for c in (1, 5, 7)] == ["1.3", "a2", "Revise & Resubmit"]
    none = wb["(No vendor)"]
    rows = [[none.cell(r, c).value for c in (1, 4, 5, 7, 9)] for r in range(6, none.max_row + 1)]
    assert rows[0][:4] == ["1.2", "A", None, "Overdue"] and rows[0][4] > 2000
    assert [r[0] for r in rows] == ["1.2", "1.3", "2.1", "2.2", "2.3"]
    assert rows[-1][3:] == ["Not Submitted", None]
    assert wb.properties.creator == "docbase (system generated)"


def test_received_and_outstanding_in_list_and_summary(api_client, project, demo):
    _attach(api_client, project, demo, "1.3/b1", "b1.pdf")
    items = _by_key(_call(api_client, "GET", project, "deliverables")["items"], demo)
    assert (items["1.3/b1"]["received"], items["1.3/b1"]["outstanding"]) == (True, False)
    assert (items["2.3"]["received"], items["2.3"]["outstanding"]) == (False, True)
    assert items["1.1"]["outstanding"] is False  # 已批准的没文件也不算缺件
    overall = _call(api_client, "GET", project, "deliverables/summary")["overall"]
    # 应交 10（11 减不适用 1.2/B），收到 1；缺件 = 没文件且没批准、不是不适用的 7 项
    assert (overall["received"], overall["applicable"], overall["outstanding"]) == (1, 10, 7)
    assert overall["received_rate"] == 0.1
    page = _call(api_client, "GET", project, "deliverables?received=false&outstanding=true")
    assert "1.3/b1" not in _by_key(page["items"], demo) and page["total"] == 7


def test_summary_ignores_card_filters(api_client, project, demo):
    """点卡片 / 按钮加的筛选（状态、责任方、逾期、收集情况）不改变顶部数字；范围筛选会。"""
    whole = _call(api_client, "GET", project, "deliverables/summary")["overall"]
    for q in ("responsible=vendor&status=not_submitted", "overdue=true", "received=false", "outstanding=true",
              "no_evidence=true"):
        assert _call(api_client, "GET", project, f"deliverables/summary?{q}")["overall"] == whole, q
    scoped = _call(api_client, "GET", project, f"deliverables/summary?milestone_id={demo['m2']['id']}")["overall"]
    assert scoped["total"] == 3


def test_outstanding_export_filters_and_chinese(api_client, project, demo):
    resp = _call(api_client, "GET", project, "deliverables/outstanding.xlsx?vendor=V1&lang=zh")
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["汇总", "V1"]
    assert wb["V1"].cell(6, 7).value == "需修改重交"


# ---------------------------------------------------------------- MDL 导出和矩阵口径一致


def test_mdl_export_uses_deliverable_due_date(session, project, demo):
    data = mx.export_current(session, project["id"], today=date(2026, 10, 24))
    wb = load_workbook(io.BytesIO(data))
    ws = wb["MDL"]
    row = next(r for r in range(7, ws.max_row + 1) if ws.cell(r, mx.C_NO).value == "2.2")
    assert ws.cell(row, mx.C_STATUS).value == "Overdue"
    assert ws.cell(row, mx.C_DUE).value.date() == date(2020, 2, 1)
