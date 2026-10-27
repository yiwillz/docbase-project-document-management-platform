import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app.services.structure_template import build_template, build_workbook, parse_template

USER = {"X-User": "alice"}


def _fill(milestones=(), document_types=(), partitions=(), lang="zh", mutate=None) -> bytes:
    """在空白模板里按列顺序填数据。"""
    wb = build_workbook(lang)
    for ws, rows in zip(wb.worksheets[:3], (milestones, document_types, partitions), strict=True):
        for row in rows:
            ws.append(list(row))
    if mutate:
        mutate(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# 里程碑：名称、目标、说明、备注、截止日期
MILESTONES = [("Kickoff", None, None, None, None),
              ("Design", "Freeze design", None, None, date(2026, 6, 7))]
# 文件种类：里程碑、编号、名称、分类、按设备、模块、截止日期、应含内容…
DOCS = [
    ("M1", "1.1", "Schedule", "Control"),
    ("M2", "2.1", "Spec", "Design", "是", "Automation, Alpha"),
    ("2", None, "Layout", "Design", "否"),
    ("M2", "2.3", "P&ID", "Design", "Y", "Alpha", None, "Lines"),
]
PARTS = [("2.1", "Alpha", "Approved", "ok"), ("2.3", "Thermal", "Pending", None)]


def test_template_has_three_sheets_and_guide():
    wb = load_workbook(io.BytesIO(build_template()))
    assert wb.sheetnames == ["里程碑", "文件种类", "模块分区", "说明"]
    assert [c.value for c in wb["里程碑"][1]] == ["名称", "目标", "说明", "备注", "截止日期"]
    en = load_workbook(io.BytesIO(build_template("en")))
    assert en.sheetnames[:3] == ["Milestones", "Document Types", "Module Partitions"]


def test_parse_numbers_modules_and_partitions():
    parsed = parse_template(_fill(MILESTONES, DOCS, PARTS))
    assert parsed.ok, parsed.errors
    assert [d["submittal_no"] for d in parsed.document_types] == ["1.1", "2.1", "2.2", "2.3"]
    spec, layout, pid = parsed.document_types[1:]
    assert spec["per_equipment"] is True and layout["per_equipment"] is False
    assert spec["modules"] == ["Automation", "Alpha"]
    assert pid["modules"] == ["Alpha", "Thermal"]  # 分区表里的模块自动补上
    assert spec["partitions"]["Alpha"] == {"mdl_status_text": "Approved", "mdl_remarks": "ok"}
    assert parsed.modules == ["Automation", "Alpha", "Thermal"]
    assert parsed.milestones[1]["due_date"] == date(2026, 6, 7)
    assert parsed.summary() == {"milestones": 2, "document_types": 4, "per_equipment": 2,
                                "per_module": 2, "modules": 3, "partitions": 4}


def test_english_headers_accepted():
    parsed = parse_template(_fill(MILESTONES, DOCS, PARTS, lang="en"))
    assert parsed.ok and len(parsed.document_types) == 4


@pytest.mark.parametrize("docs, parts, code", [
    ([("M3", None, "X")], [], "template.unknown_milestone"),
    ([("M1", "1.2", "X")], [], "template.number_mismatch"),     # 第一行只能是 1.1
    ([("M1", None, None)], [], "template.required"),
    ([("M1", None, "X", None, "maybe")], [], "template.bad_value"),
    ([("M1", None, "X", None, None, None, "next week")], [], "template.bad_value"),
    ([("M1", None, "X")], [("9.9", "Alpha")], "template.unknown_document_type"),
    ([("M1", None, "X")], [("1.1", "Alpha"), ("1.1", "Alpha")], "template.duplicate_partition"),
])
def test_parse_errors(docs, parts, code):
    parsed = parse_template(_fill(MILESTONES, docs, parts))
    assert code in [e.code for e in parsed.errors]
    error = next(e for e in parsed.errors if e.code == code)
    assert error.sheet and error.row  # 指出哪张表、第几行


def test_float_number_that_lost_its_zero_is_caught():
    """Excel 把 4.10 存成数字时会变成 4.1，核对编号能发现。"""
    milestones = [(f"M{i}",) for i in range(1, 5)]
    docs = [("M4", f"4.{i}", f"D{i}") for i in range(1, 10)] + [("M4", 4.1, "D10")]
    parsed = parse_template(_fill(milestones, docs))
    [error] = parsed.errors
    assert error.code == "template.number_mismatch" and "4.10" in error.message


def test_missing_sheet_and_column():
    def drop(wb):
        del wb["文件种类"]
    parsed = parse_template(_fill(MILESTONES, mutate=drop))
    assert [e.code for e in parsed.errors] == ["template.missing_sheet"]

    def rename(wb):
        wb["文件种类"]["C1"] = "标题"
    parsed = parse_template(_fill(MILESTONES, DOCS, mutate=rename))
    codes = [e.code for e in parsed.errors]
    assert "template.missing_column" in codes
    assert "template.unknown_column" in [w.code for w in parsed.warnings]


def test_not_an_excel_file():
    assert parse_template(b"hello").errors[0].code == "template.unreadable"


# ---------------------------------------------------------------- 接口：下载 → 预览 → 确认


@pytest.fixture
def project(api_client):
    return api_client.post("/api/projects", json={"code": "TPL-1", "name": "x"},
                           headers=USER).json()


def _preview(api_client, project, data):
    return api_client.post(f"/api/projects/{project['id']}/structure-import/preview",
                           files={"file": ("s.xlsx", data)}, headers=USER)


def test_download_template(api_client):
    resp = api_client.get("/api/structure-template")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert load_workbook(io.BytesIO(resp.content)).sheetnames[0] == "里程碑"


def test_preview_then_confirm(api_client, project):
    resp = _preview(api_client, project, _fill(MILESTONES, DOCS, PARTS))
    assert resp.status_code == 200, resp.text
    preview = resp.json()
    assert preview["ok"] is True
    assert preview["milestones"][1]["document_types"][0] == {
        "submittal_no": "2.1", "name": "Spec", "per_equipment": True,
        "modules": ["Automation", "Alpha"]}
    # 预览不写入
    tree = api_client.get(f"/api/projects/{project['id']}/structure").json()
    assert tree["milestones"] == []

    resp = api_client.post(f"/api/projects/{project['id']}/structure-import/confirm",
                           json={"token": preview["token"]}, headers=USER)
    assert resp.status_code == 200, resp.text
    assert resp.json() == {"milestones": 2, "modules": 3, "document_types": 4, "partitions": 4}

    tree = api_client.get(f"/api/projects/{project['id']}/structure").json()
    assert [m["code"] for m in tree["milestones"]] == ["M1", "M2"]
    spec = tree["milestones"][1]["document_types"][0]
    assert (spec["submittal_no"], spec["per_equipment"], spec["per_module"]) == ("2.1", True, True)
    alpha = next(p for p in spec["partitions"] if p["module_name"] == "Alpha")
    assert (alpha["mdl_status_text"], alpha["mdl_remarks"]) == ("Approved", "ok")

    # 同一个 token 不能导两次；项目已有结构也不能再导
    resp = api_client.post(f"/api/projects/{project['id']}/structure-import/confirm",
                           json={"token": preview["token"]}, headers=USER)
    assert resp.json()["detail"]["code"] == "structure_import.already_imported"
    again = _preview(api_client, project, _fill(MILESTONES, DOCS)).json()
    resp = api_client.post(f"/api/projects/{project['id']}/structure-import/confirm",
                           json={"token": again["token"]}, headers=USER)
    assert resp.json()["detail"]["code"] == "structure.target_not_empty"


def test_confirm_refused_when_preview_has_errors(api_client, project):
    preview = _preview(api_client, project, _fill(MILESTONES, [("M9", None, "X")])).json()
    assert preview["ok"] is False and preview["errors"][0]["row"] == 2
    resp = api_client.post(f"/api/projects/{project['id']}/structure-import/confirm",
                           json={"token": preview["token"]}, headers=USER)
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "structure_import.has_errors"


def test_token_belongs_to_its_project(api_client, project):
    preview = _preview(api_client, project, _fill(MILESTONES, DOCS)).json()
    other = api_client.post("/api/projects", json={"code": "TPL-2", "name": "y"},
                            headers=USER).json()
    resp = api_client.post(f"/api/projects/{other['id']}/structure-import/confirm",
                           json={"token": preview["token"]}, headers=USER)
    assert resp.status_code == 404
