import io
import json
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import Deliverable, File, FileDeliverable, FileSource
from app.services.equipment_import import parse_equipment

USER = {"X-User": "alice"}


def _xlsx(rows, *, title="Equipment", before=()) -> bytes:
    """造一份虚构的设备清单：before 是表头上方的杂项行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for row in list(before) + list(rows):
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


CN = _xlsx([
    ("设备编号", "设备名称", "类别", "供应商", "台数", "别名", "功率(kW)", "进场日期"),
    (71, "Alpha bench A", "种类A", "Vendor A", 2, "WB-A, 设备A", 15.5, date(2026, 5, 1)),
    (72, "Alpha bench B", "种类A", "Vendor A", None, None, 15, None),
    ("80", "Dryer", "种类B", "Vendor B", 1, None, None, None),
], before=[("项目设备清单（虚构）",), ()])

EN = _xlsx([
    ("Qty", "Name", "Tag", "Supplier", "Category", "Voltage"),
    (1, "Robot", "A1", "Vendor C", "Automation", "380V"),
])


@pytest.fixture
def project(api_client):
    return api_client.post("/api/projects", json={"code": "EQ-1", "name": "equipment"},
                           headers=USER).json()


def _url(project, path=""):
    return f"/api/projects/{project['id']}/{path}"


def _call(api_client, method, project, path, body=None, status=200):
    resp = api_client.request(method, _url(project, path), json=body, headers=USER)
    assert resp.status_code == status, resp.text
    return resp.json() if resp.content else None


def _preview(api_client, project, data, **form):
    resp = api_client.post(_url(project, "equipment-import/preview"),
                           files={"file": ("list.xlsx", data)}, data=form, headers=USER)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _import(api_client, project, data, **confirm):
    preview = _preview(api_client, project, data)
    assert preview["ok"], preview["errors"]
    return _call(api_client, "POST", project, "equipment-import/confirm",
                 {"token": preview["token"], **confirm})


# ---------------------------------------------------------------- 解析


def test_parse_chinese_headers_and_extra_columns():
    parsed = parse_equipment(CN)
    assert parsed.ok and parsed.header_row == 3  # 跳过上方的标题行
    fields = {c["header"]: c["field"] for c in parsed.columns}
    assert fields["设备编号"] == "tag" and fields["类别"] == "equipment_type"
    assert fields["功率(kW)"] == "attr"
    first = parsed.rows[0]
    assert (first["tag"], first["quantity"], first["aliases"]) == ("71", 2, ["WB-A", "设备A"])
    assert first["attrs"] == {"功率(kW)": 15.5, "进场日期": "2026-05-01"}
    assert parsed.rows[1]["quantity"] is None and parsed.rows[1]["attrs"] == {"功率(kW)": 15}


def test_parse_english_headers_any_order():
    parsed = parse_equipment(EN)
    row = parsed.rows[0]
    assert (row["tag"], row["name"], row["vendor"], row["equipment_type"]) == (
        "A1", "Robot", "Vendor C", "Automation")
    assert row["attrs"] == {"Voltage": "380V"}


def test_unknown_headers_need_manual_mapping():
    data = _xlsx([("Item", "Machine", "Maker"), ("X1", "Press", "Vendor D")])
    parsed = parse_equipment(data)
    assert not parsed.ok
    assert [e.code for e in parsed.errors] == ["equipment_import.missing_field"]  # 只有设备名必填
    assert parsed.warnings[0].code == "equipment_import.header_guessed"
    assert [c["header"] for c in parsed.columns] == ["Item", "Machine", "Maker"]  # 给出全部列
    mapped = parse_equipment(data, header_row=1,
                             mapping={"A": "tag", "B": "name", "C": "vendor"})
    assert mapped.ok and mapped.rows[0]["vendor"] == "Vendor D"


@pytest.mark.parametrize("rows, code", [
    ([("设备编号", "设备名"), ("1", None)], "equipment_import.required"),
    ([("设备编号", "设备名"), (None, "零件"), (None, "零件")], "equipment_import.duplicate_tag"),
    ([("设备编号", "设备名"), ("1", "X"), ("1", "Y")], "equipment_import.duplicate_tag"),
    ([("设备编号", "设备名", "台数"), ("1", "X", "两台")], "equipment_import.bad_quantity"),
])
def test_parse_row_errors(rows, code):
    parsed = parse_equipment(_xlsx(rows))
    assert code in [e.code for e in parsed.errors]
    assert next(e for e in parsed.errors if e.code == code).row is not None


def test_template_round_trip(api_client):
    resp = api_client.get("/api/equipment-template")
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    ws.append(["9", "Furnace", "热处理", "Vendor E", 1, None])
    buffer = io.BytesIO()
    wb.save(buffer)
    parsed = parse_equipment(buffer.getvalue())
    assert parsed.ok and parsed.rows[0]["tag"] == "9"


# ---------------------------------------------------------------- 导入与比对


def test_import_then_diff(api_client, project):
    preview = _preview(api_client, project, CN)
    assert preview["diff"]["summary"]["added"] == 3
    assert preview["diff"]["new_types"] == ["种类A", "种类B"]
    counts = _call(api_client, "POST", project, "equipment-import/confirm",
                   {"token": preview["token"]})
    assert counts["added"] == 3
    types = [t["name"] for t in _call(api_client, "GET", project, "equipment-types")]
    assert set(types) == {"种类A", "种类B"}  # 设备种类自动新建

    # 第二次：改一台的名称和功率，少了一台
    changed = _xlsx([
        ("设备编号", "设备名称", "类别", "供应商", "台数", "别名", "功率(kW)"),
        (71, "Alpha bench A (rev)", "种类A", "Vendor A", 2, "WB-A, 设备A", 18),
        (72, "Alpha bench B", "种类A", "Vendor A", None, None, 15),
    ])
    preview = _preview(api_client, project, changed)
    diff = preview["diff"]
    assert diff["summary"] == {"added": 0, "modified": 1, "reactivated": 0, "missing": 1,
                               "unchanged": 1}
    [mod] = diff["modified"]
    assert mod["changes"]["name"] == ["Alpha bench A", "Alpha bench A (rev)"]
    assert mod["changes"]["attrs"][1] == {"功率(kW)": 18}  # 原来的进场日期这次没有了
    assert diff["missing"] == [{"tag": "80", "name": "Dryer"}]
    counts = _call(api_client, "POST", project, "equipment-import/confirm",
                   {"token": preview["token"]})
    assert counts == {"added": 0, "modified": 1, "reactivated": 0, "deactivated": 0}
    assert len(_call(api_client, "GET", project, "equipment")) == 3  # 没有了的只提示，不停用


def test_import_deactivate_missing_then_reactivate(api_client, project):
    _import(api_client, project, CN)
    only_71 = _xlsx([("设备编号", "设备名称"), (71, "Alpha bench A")])
    counts = _import(api_client, project, only_71, deactivate_missing=True)
    assert counts["deactivated"] == 2
    assert [e["tag"] for e in _call(api_client, "GET", project, "equipment")] == ["71"]
    preview = _preview(api_client, project, CN)
    assert len(preview["diff"]["reactivated"]) == 2
    counts = _call(api_client, "POST", project, "equipment-import/confirm",
                   {"token": preview["token"]})
    assert counts["reactivated"] == 2


def test_confirm_refused_with_errors_or_twice(api_client, project):
    bad = _preview(api_client, project, _xlsx([("Item", "Machine"), ("1", "X")]))
    resp = api_client.post(_url(project, "equipment-import/confirm"),
                           json={"token": bad["token"]}, headers=USER)
    assert resp.json()["detail"]["code"] == "equipment_import.has_errors"
    good = _preview(api_client, project, EN)
    _call(api_client, "POST", project, "equipment-import/confirm", {"token": good["token"]})
    resp = api_client.post(_url(project, "equipment-import/confirm"),
                           json={"token": good["token"]}, headers=USER)
    assert resp.json()["detail"]["code"] == "equipment_import.already_imported"


def test_mapping_via_api(api_client, project):
    data = _xlsx([("Item", "Machine", "Note"), ("X1", "Press", "old")])
    preview = _preview(api_client, project, data, header_row="1",
                       mapping=json.dumps({"A": "tag", "B": "name", "C": "ignore"}))
    assert preview["ok"] and preview["diff"]["summary"]["added"] == 1


# ---------------------------------------------------------------- 设备管理


def test_tag_unique_and_not_reused(api_client, project):
    item = _call(api_client, "POST", project, "equipment", {"tag": "72", "name": "Mixer"}, 201)
    _call(api_client, "POST", project, f"equipment/{item['id']}/deactivate")
    resp = api_client.post(_url(project, "equipment"), json={"tag": "72", "name": "New"},
                           headers=USER)
    assert resp.json()["detail"]["code"] == "equipment.tag_taken"


def test_units_only_for_serial_numbers(api_client, project, session):
    """单台记录不再按台数自动建：只有填了出厂编号的台才有记录（台数几千也不会建几千条）。"""
    big = _call(api_client, "POST", project, "equipment", {"name": "Beta part", "quantity": 3000}, 201)
    assert _call(api_client, "GET", project, f"equipment/{big['id']}/units") == []
    item = _call(api_client, "POST", project, "equipment",
                 {"tag": "72", "name": "Mixer", "quantity": 3}, 201)
    assert _call(api_client, "GET", project, f"equipment/{item['id']}/units") == []
    _call(api_client, "PUT", project, f"equipment/{item['id']}/units",
          {"serials": [{"seq": 1, "serial_no": "SN-1"}, {"seq": 3, "serial_no": "SN-3"},
                       {"seq": 2, "serial_no": "  "}]})
    units = _call(api_client, "GET", project, f"equipment/{item['id']}/units")
    assert [(u["seq"], u["serial_no"]) for u in units] == [(1, "SN-1"), (3, "SN-3")]
    resp = api_client.put(_url(project, f"equipment/{item['id']}/units"), headers=USER,
                          json={"serials": [{"seq": 4, "serial_no": "SN-4"}]})
    assert resp.status_code == 404  # 超过台数
    _call(api_client, "PATCH", project, f"equipment/{item['id']}", {"quantity": 2})
    assert [u["seq"] for u in _call(api_client, "GET", project, f"equipment/{item['id']}/units")] == [1]
    _call(api_client, "PATCH", project, f"equipment/{item['id']}", {"quantity": 3})
    units = _call(api_client, "GET", project, f"equipment/{item['id']}/units")
    assert units[1]["serial_no"] == "SN-3"  # 台数减少时只停用，加回来出厂编号还在


def test_untagged_equipment_identified_by_name(api_client, project):
    """没编号的设备用名称区分：没编号的名称不能重复；有编号的可以和它同名。"""
    basket = _call(api_client, "POST", project, "equipment", {"name": "Alpha part", "quantity": 2000}, 201)
    assert basket["tag"] is None
    resp = api_client.post(_url(project, "equipment"), json={"name": "Alpha part"}, headers=USER)
    assert resp.status_code == 409 and resp.json()["detail"]["code"] == "equipment.name_taken"
    _call(api_client, "POST", project, "equipment", {"tag": "W-1", "name": "Alpha part"}, 201)
    # 去掉编号：变成没编号，名称和已有的没编号设备冲突
    tagged = _call(api_client, "POST", project, "equipment", {"tag": "W-2", "name": "Rack"}, 201)
    _call(api_client, "PATCH", project, f"equipment/{tagged['id']}", {"tag": ""})
    resp = api_client.post(_url(project, "equipment"), json={"name": "Rack"}, headers=USER)
    assert resp.json()["detail"]["code"] == "equipment.name_taken"


def test_import_matches_untagged_by_name(api_client, project):
    first = _xlsx([("设备编号", "设备名", "台数"), ("72", "Mixer", None), (None, "Beta part", 3000)])
    _import(api_client, project, first)
    second = _xlsx([("设备编号", "设备名", "台数"), ("72", "Mixer", None), (None, "Beta part", 3500)])
    diff = _preview(api_client, project, second)["diff"]
    assert diff["summary"]["added"] == 0 and diff["summary"]["modified"] == 1
    assert diff["modified"][0]["changes"] == {"quantity": [3000, 3500]}


def test_history_and_restore(api_client, project):
    item = _call(api_client, "POST", project, "equipment",
                 {"tag": "72", "name": "Mixer", "vendor": "Vendor A"}, 201)
    _call(api_client, "PATCH", project, f"equipment/{item['id']}",
          {"name": "Mixer v2", "vendor": "Vendor B", "reason": "供应商变更"})
    _call(api_client, "PATCH", project, f"equipment/{item['id']}", {"quantity": 2})
    history = _call(api_client, "GET", project, f"equipment/{item['id']}/history")
    assert [h["action"] for h in history] == ["update", "update", "create"]
    rename = history[1]
    assert rename["changes"]["name"] == ["Mixer", "Mixer v2"] and rename["reason"] == "供应商变更"

    restored = _call(api_client, "POST", project, f"equipment/{item['id']}/restore",
                     {"change_log_id": rename["id"]})
    assert (restored["name"], restored["vendor"], restored["quantity"]) == ("Mixer", "Vendor A", 2)
    latest = _call(api_client, "GET", project, f"equipment/{item['id']}/history")[0]
    assert latest["action"] == "restore" and f"#{rename['id']}" in latest["reason"]
    resp = api_client.post(_url(project, f"equipment/{item['id']}/restore"),
                           json={"change_log_id": history[-1]["id"]}, headers=USER)
    assert resp.json()["detail"]["code"] == "equipment.restore_not_update"


def test_tag_locked_once_files_exist(api_client, project, session):
    m1 = _call(api_client, "POST", project, "milestones", {"name": "Design"}, 201)
    [doc] = _call(api_client, "POST", project, "document-types/bulk",
                  {"milestone_id": m1["id"], "rows": [{"name": "Cut sheet"}]}, 201)
    item = _call(api_client, "POST", project, "equipment", {"tag": "72", "name": "Mixer"}, 201)
    _call(api_client, "POST", project, "document-types/bulk-equipment",
          {"document_type_ids": [doc["id"]], "equipment_ids": [item["id"]], "mode": "add"})
    _call(api_client, "PATCH", project, f"equipment/{item['id']}", {"tag": "73"})  # 没文件，能改
    deliverable = session.scalars(select(Deliverable).where(
        Deliverable.equipment_id == item["id"])).one()
    file = File(project_id=project["id"], source=FileSource.VENDOR)
    session.add(file)
    session.flush()
    session.add(FileDeliverable(file_id=file.id, deliverable_id=deliverable.id))
    session.commit()
    resp = api_client.patch(_url(project, f"equipment/{item['id']}"), json={"tag": "74"},
                            headers=USER)
    assert resp.json()["detail"]["code"] == "equipment.tag_locked"


def test_equipment_changes_update_deliverables(api_client, project, session):
    m1 = _call(api_client, "POST", project, "milestones", {"name": "Design"}, 201)
    [doc] = _call(api_client, "POST", project, "document-types/bulk",
                  {"milestone_id": m1["id"], "rows": [{"name": "Cut sheet"}]}, 201)
    _import(api_client, project, CN)
    _call(api_client, "POST", project, "document-types/bulk-equipment",
          {"document_type_ids": [doc["id"]], "all_equipment": True, "mode": "add"})
    summary = _call(api_client, "GET", project, "deliverables/summary")
    assert summary["by_kind"] == {"per_equipment": 3}
    dryer = next(e for e in _call(api_client, "GET", project, "equipment") if e["tag"] == "80")
    _call(api_client, "POST", project, f"equipment/{dryer['id']}/deactivate")
    summary = _call(api_client, "GET", project, "deliverables/summary")
    # 作废 2 个：改成按设备时原来"整个种类"的 1 个 + 停用设备的 1 个
    assert (summary["active"], summary["obsolete"]) == (2, 2)


def test_module_equipment_types(api_client, project):
    _import(api_client, project, CN)
    alpha = _call(api_client, "POST", project, "modules", {"name": "Alpha"}, 201)
    types = {t["name"]: t["id"] for t in _call(api_client, "GET", project, "equipment-types")}
    result = _call(api_client, "PUT", project, f"modules/{alpha['id']}/equipment-types",
                   {"equipment_type_ids": [types["种类A"], types["种类B"]]})
    assert {t["name"] for t in result} == {"种类A", "种类B"}
    result = _call(api_client, "PUT", project, f"modules/{alpha['id']}/equipment-types",
                   {"equipment_type_ids": []})
    assert result == []  # 可以不对应


def test_pasted_rows_preview_and_confirm(api_client, project):
    """页面上粘贴多行设备：和 Excel 导入同样校验、比对，一次确认。"""
    body = {"rows": [{"tag": "72", "name": "Mixer", "equipment_type": "Alpha", "quantity": "2"},
                     {"name": "Beta part", "equipment_type": "Part", "quantity": 3000},
                     {}]}
    preview = _call(api_client, "POST", project, "equipment-import/preview-rows", body)
    assert preview["ok"] and preview["diff"]["summary"]["added"] == 2
    assert preview["diff"]["new_types"] == ["Alpha", "Part"]
    counts = _call(api_client, "POST", project, "equipment-import/confirm", {"token": preview["token"]})
    assert counts["added"] == 2
    items = _call(api_client, "GET", project, "equipment")
    assert sorted((e["tag"], e["name"], e["quantity"]) for e in items if e["tag"] is None) == [
        (None, "Beta part", 3000)]
    bad = _call(api_client, "POST", project, "equipment-import/preview-rows",
                {"rows": [{"tag": "9", "quantity": "x"}, {"name": "A", "quantity": "two"}]})
    assert [e["code"] for e in bad["errors"]] == ["equipment_import.required", "equipment_import.bad_quantity"]
