"""第 8B 步：闭环检查、里程碑关闭、客户新增要求的接受情况、合同外新增要求清单、清理掉的字段。

用 test_transmittals 的虚构项目 TX-1：
  1.1 Spec（合同规定，依据 App D 3.1）、1.2 Layout 按模块 A、B（合同规定，没写依据）、
  1.3 Noise study（客户要求，依据 CL-015，待确认）、1.4 Checklist（内部要求）
"""

import io

from openpyxl import load_workbook

from app.services import mdl_export as mx
from tests.test_transmittals import USER, _url, call, demo, items, project, submit  # noqa: F401


def _check(api_client, project, demo):
    return call(api_client, "GET", project, f"milestones/{demo['m1']['id']}/closeout")


def _issues(result):
    return {(r["submittal_no"], r["module_name"]): r["issues"] for r in result["items"]}


def _approve(api_client, project, ids, code="1", date_="2026-11-01", ref="CL-100"):
    t = submit(api_client, project, ids, submitted_date="2026-10-20")
    call(api_client, "POST", project, f"transmittals/{t['id']}/reply",
         {"code": code, "reply_date": date_, "reply_ref": ref})
    return t


# ---------------------------------------------------------------- 闭环检查


def test_check_lists_what_blocks_closing(api_client, project, demo):
    result = _check(api_client, project, demo)
    assert (result["closable"], result["blocking"]) == (False, 4)
    assert _issues(result) == {
        ("1.1", None): ["not_done"], ("1.2", "A"): ["not_done"], ("1.2", "B"): ["not_done"],
        ("1.3", None): ["not_done", "requirement_unconfirmed"],
        ("1.4", None): ["internal_open"]}  # 内部要求只提示
    assert [d["submittal_no"] for d in result["documents"]] == ["1.2"]  # 没写要求依据
    assert result["counts"]["no_requirement_basis"] == 1 and result["warnings"] == 2


def test_closable_after_final_approvals_with_evidence(api_client, project, demo):
    d = demo["d"]
    _approve(api_client, project, [d["1.1"]])
    _approve(api_client, project, [d["1.2/A"]], code="2")  # 带意见批准
    call(api_client, "PATCH", project, f"deliverables/{d['1.2/B']}", {"status": "not_applicable"})
    call(api_client, "PATCH", project, f"document-types/{demo['docs'][2]['id']}",
         {"acceptance": "accepted", "acceptance_ref": "Reply RL-03"})
    call(api_client, "PATCH", project, f"deliverables/{d['1.3']}", {"status": "approved"})
    issues = _issues(_check(api_client, project, demo))
    assert issues[("1.1", None)] == []
    assert issues[("1.2", "A")] == ["not_final"]
    assert issues[("1.2", "B")] == ["no_evidence"]  # 不适用没写依据
    assert issues[("1.3", None)] == ["no_evidence"]  # 手改的批准，没有客户回复
    _approve(api_client, project, [d["1.2/A"], d["1.3"]], ref="CL-101")
    call(api_client, "PATCH", project, f"deliverables/{d['1.2/B']}",
         {"status_basis": "Client letter CL-099: module B deleted"})
    result = _check(api_client, project, demo)
    assert (result["closable"], result["blocking"]) == (True, 0)
    assert result["stats"]["approved"] == 3


# ---------------------------------------------------------------- 关闭、重新打开


def test_close_force_reopen_and_change_after_close(api_client, project, demo):
    mid = demo["m1"]["id"]
    resp = api_client.post(_url(project, f"milestones/{mid}/close"), headers=USER, json={})
    assert resp.status_code == 409 and resp.json()["detail"]["code"] == "milestone.not_closable"
    resp = api_client.post(_url(project, f"milestones/{mid}/close"), headers=USER, json={"force": True})
    assert resp.json()["detail"]["code"] == "milestone.force_needs_reason"
    ms = call(api_client, "POST", project, f"milestones/{mid}/close", {
        "force": True, "reason": "客户书面同意带尾项关闭", "basis": "PA-03",
        "closed_date": "2026-11-30"})
    assert (ms["closed"], ms["closed_date"], ms["closed_by"]) == (True, "2026-11-30", "alice")
    assert ms["close_basis"] == "PA-03（带 4 项未闭环关闭：客户书面同意带尾项关闭）"
    assert ms["changed_after_close"] is False
    call(api_client, "PATCH", project, f"deliverables/{demo['d']['1.1']}", {"remarks": "after close"})
    assert _check(api_client, project, demo)["milestone"]["changed_after_close"] is True
    resp = api_client.post(_url(project, f"milestones/{mid}/close"), headers=USER, json={"force": True, "reason": "x"})
    assert resp.json()["detail"]["code"] == "milestone.closed"
    resp = api_client.post(_url(project, f"milestones/{mid}/reopen"), headers=USER, json={"reason": ""})
    assert resp.status_code == 422
    ms = call(api_client, "POST", project, f"milestones/{mid}/reopen", {"reason": "补交 1.2"})
    assert (ms["closed"], ms["close_basis"]) == (False, None)


def test_closeout_export(api_client, project, demo):
    _approve(api_client, project, [demo["d"]["1.1"]], ref="CL-100")
    url = f"milestones/{demo['m1']['id']}/closeout/export"
    wb = load_workbook(io.BytesIO(call(api_client, "GET", project, url).content))
    ws = wb.worksheets[0]
    assert ws["A1"].value == "MILESTONE CLOSEOUT LIST" and ws["B5"].value == "3 item(s) open"
    first = [ws.cell(9, c).value for c in (1, 5, 6, 8, 9)]
    assert first == ["1.1", "Approved", "1", "CL-100", "TX-1-TR-0001"]
    assert ws.cell(10, 11).value == "Not approved"
    zh = load_workbook(io.BytesIO(call(api_client, "GET", project, url, params={"lang": "zh"}).content))
    assert zh.worksheets[0]["A1"].value == "里程碑关闭清单"
    for lang in ("en", "zh"):
        pdf = call(api_client, "GET", project, url, params={"format": "pdf", "lang": lang}).content
        assert pdf.startswith(b"%PDF")


# ---------------------------------------------------------------- 客户新增要求


def test_acceptance_rules(api_client, project, demo):
    docs = demo["docs"]
    assert (docs[2]["acceptance"], docs[0]["acceptance"]) == ("pending", None)
    resp = api_client.patch(_url(project, f"document-types/{docs[0]['id']}"), headers=USER,
                            json={"acceptance": "accepted"})
    assert resp.json()["detail"]["code"] == "document_type.acceptance_not_client"
    resp = api_client.patch(_url(project, f"document-types/{docs[2]['id']}"), headers=USER,
                            json={"acceptance": "declined"})
    assert resp.json()["detail"]["code"] == "document_type.decline_needs_ref"
    doc = call(api_client, "PATCH", project, f"document-types/{docs[2]['id']}",
               {"acceptance": "declined", "acceptance_ref": "RL-07", "variation_ref": "VO-2"})
    assert (doc["acceptance"], doc["variation_ref"]) == ("declined", "VO-2")
    it = items(api_client, project)[demo["d"]["1.3"]]
    assert (it["status"], it["status_basis"], it["no_evidence"]) == (
        "not_applicable", "客户要求不接受：RL-07", False)
    assert _issues(_check(api_client, project, demo))[("1.3", None)] == []  # 不挡关闭
    # 客户 MDL 里留着：不适用 + 拒绝依据
    wb = load_workbook(io.BytesIO(call(api_client, "GET", project, "mdl/export").content))
    ws = wb["MDL"]
    row = next(r for r in range(7, ws.max_row + 1) if ws.cell(r, mx.C_NO).value == "1.3")
    assert (ws.cell(row, mx.C_STATUS).value, ws.cell(row, mx.C_REM).value) == (
        "N/A", "Not accepted – RL-07")
    # 改回接受：回到未提交
    call(api_client, "PATCH", project, f"document-types/{docs[2]['id']}", {"acceptance": "accepted"})
    assert items(api_client, project)[demo["d"]["1.3"]]["status"] == "not_submitted"
    # 改成合同规定：接受情况清空
    doc = call(api_client, "PATCH", project, f"document-types/{docs[2]['id']}", {"origin": "contract"})
    assert doc["acceptance"] is None


def test_new_deliverables_of_declined_requirement_are_na(api_client, project, demo):
    doc = call(api_client, "POST", project, "document-types/bulk", {
        "milestone_id": demo["m1"]["id"], "rows": [{
            "name": "Extra report", "origin": "client", "acceptance": "declined",
            "acceptance_ref": "RL-08"}]}, 201)[0]
    rows = items(api_client, project, document_type_id=doc["id"])
    assert [r["status"] for r in rows.values()] == ["not_applicable"]


def test_summary_by_origin_and_client_requirements_report(api_client, project, demo):
    _approve(api_client, project, [demo["d"]["1.3"]])
    s = call(api_client, "GET", project, "deliverables/summary")
    assert {o: (v["total"], v["approved"]) for o, v in s["by_origin"].items()} == {
        "client": (1, 1), "contract": (3, 0), "internal": (1, 0)}
    call(api_client, "PATCH", project, f"document-types/{demo['docs'][2]['id']}",
         {"acceptance": "accepted", "acceptance_ref": "RL-03", "variation_ref": "VO-1"})
    wb = load_workbook(io.BytesIO(call(api_client, "GET", project, "client-requirements.xlsx").content))
    ws = wb.worksheets[0]
    assert [ws.cell(6, c).value for c in (1, 2, 3, 5, 6, 7, 8, 11, 12, 13)] == [
        "1.3", "Noise study", "M1", "CL-015", "Accepted", "RL-03", "VO-1", 1, 1, 1]
    assert ws.cell(6, 4).value.date().isoformat() == "2026-09-01"
    assert ws.max_row == 6  # 只有客户要求


# ---------------------------------------------------------------- 清理掉的字段


def test_removed_fields_are_gone(api_client, project, demo):
    doc = demo["docs"][0]
    for field in ("epc_reference", "doc_count", "knowledge_form", "visibility"):
        assert field not in doc
    assert "default_mdl_template_id" not in project and project["review_days"] == 14
