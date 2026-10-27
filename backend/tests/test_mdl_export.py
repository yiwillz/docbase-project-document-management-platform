import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models import (
    Deliverable,
    DeliverableStatus,
    Equipment,
    EquipmentType,
    Event,
    ModuleEquipmentType,
)
from app.services import mdl_export as mx
from app.services.deliverables import update_deliverable
from app.services.history import Actor

USER = {"X-User": "alice"}


@pytest.fixture
def project(api_client):
    return api_client.post("/api/projects", json={
        "code": "MX-1", "name": "Export demo", "client_name": "Client Co.",
        "contractor_name": "Contractor Co.", "mdl_doc_no": "MX-1-MDL-0001",
        "mdl_prepared_by": "Doc Controller"}, headers=USER).json()


def _url(project, path=""):
    return f"/api/projects/{project['id']}/{path}"


def _call(api_client, method, project, path, body=None, status=200):
    resp = api_client.request(method, _url(project, path), json=body, headers=USER)
    assert resp.status_code == status, resp.text
    return resp.json() if resp.content else None


@pytest.fixture
def setup(api_client, project, session):
    """M1（截止 2026-06-07）：1.1 整个种类（MDL 已批准）、1.2 按模块（A、B）、
    1.3 按设备（Alpha 模块 2 台 + 1 台其他）、1.4 内部要求；M2（无截止）：2.1 整个种类。"""
    m1 = _call(api_client, "POST", project, "milestones", {"name": "Design", "due_date": "2026-06-07"}, 201)
    m2 = _call(api_client, "POST", project, "milestones", {"name": "FAT"}, 201)
    d11, d12, d13, d14 = _call(api_client, "POST", project, "document-types/bulk", {
        "milestone_id": m1["id"], "rows": [
            {"name": "Spec & Data", "category": "Design", "mdl_status_text": "Approved",
             "mdl_remarks": "client note"},
            {"name": "Layout"}, {"name": "Cut sheet"},
            {"name": "Internal checklist", "origin": "internal"}]}, 201)
    [d21] = _call(api_client, "POST", project, "document-types/bulk",
                  {"milestone_id": m2["id"], "rows": [{"name": "FAT report"}]}, 201)
    _call(api_client, "PUT", project, f"document-types/{d12['id']}/modules", {"module_names": ["A", "B"]})
    _call(api_client, "PUT", project, f"document-types/{d13['id']}/modules", {"module_names": ["Alpha"]})
    alpha_type = EquipmentType(project_id=project["id"], name="Alpha")
    beta_type = EquipmentType(project_id=project["id"], name="Beta")
    session.add_all([alpha_type, beta_type])
    session.flush()
    eqs = [Equipment(project_id=project["id"], equipment_type_id=alpha_type.id, tag="71", name="W1", sort_order=1),
           Equipment(project_id=project["id"], equipment_type_id=alpha_type.id, tag="72", name="W2", sort_order=2),
           Equipment(project_id=project["id"], equipment_type_id=beta_type.id, tag="80", name="D1", sort_order=3)]
    session.add_all(eqs)
    session.flush()
    alpha = next(m for m in _call(api_client, "GET", project, "modules") if m["name"] == "Alpha")
    session.add(ModuleEquipmentType(module_id=alpha["id"], equipment_type_id=alpha_type.id))
    session.commit()
    _call(api_client, "POST", project, "document-types/bulk-equipment",
          {"document_type_ids": [d13["id"]], "all_equipment": True, "mode": "add"})
    w1 = session.scalars(select(Deliverable).where(Deliverable.equipment_id == eqs[0].id)).one()
    w1.status = DeliverableStatus.APPROVED
    session.commit()
    return {"m1": m1, "docs": {"1.1": d11, "1.2": d12, "1.3": d13, "1.4": d14, "2.1": d21}}


def _book(data: bytes, data_only=True):
    wb = load_workbook(io.BytesIO(data), data_only=data_only)
    meta = {r[0].value: r[1].value for r in wb["_meta"].iter_rows(max_col=2)}
    return wb, wb["MDL"], meta


def _rows(ws, meta):
    return range(int(meta["header_row"]) + 1, int(meta["last_row"]) + 1)


def _find(ws, meta, no=None, sub=None, after=None):
    """按编号找行；sub 为模块子行的名字（在编号 after 的下面）。"""
    current = None
    for r in _rows(ws, meta):
        n, title = ws.cell(r, mx.C_NO).value, (ws.cell(r, mx.C_TITLE).value or "").strip()
        current = n or current
        if no and n == no:
            return r
        if sub and not n and title == sub and current == after:
            return r
    raise AssertionError((no, sub, after))


# ---------------------------------------------------------------- Excel


def test_excel_layout(api_client, project, setup):
    resp = api_client.get(_url(project, "mdl/export"))
    assert resp.status_code == 200
    wb, ws, meta = _book(resp.content)
    assert wb.sheetnames == ["MDL", "Revision History", "Outstanding Items", "_meta"]
    assert wb["_meta"].sheet_state == "hidden" and meta["format"] == "docbase-mdl"
    assert wb.properties.creator == "docbase (system generated)"
    # 表头区：项目设置里的客户、总包、文件编号；草稿标记
    assert ws["A1"].value == "MASTER DOCUMENT LIST"
    assert (ws["B3"].value, ws["B4"].value) == ("Client Co.", "Contractor Co.")
    assert (ws["H2"].value, ws["H3"].value) == ("MX-1-MDL-0001", "DRAFT (not issued)")
    assert "System-generated" in ws.oddFooter.left.text

    numbers = [n for r in _rows(ws, meta) if (n := ws.cell(r, mx.C_NO).value) and n[0].isdigit()]
    assert numbers == ["1.1", "1.2", "1.3", "2.1"]  # 内部要求不出现
    r11 = _find(ws, meta, "1.1")
    assert (ws.cell(r11, mx.C_TITLE).value, ws.cell(r11, mx.C_STATUS).value) == ("Spec & Data", "Approved")
    assert ws.cell(r11, mx.C_REM).value == "client note"
    assert ws.cell(r11, mx.C_DUE).value.date() == date(2026, 6, 7)
    r13 = _find(ws, meta, "1.3")
    assert ws.cell(r13, mx.C_STATUS).value == "1/3 Approved"
    assert ws.cell(r13, mx.C_REM).value == "Outstanding: Alpha 1, Other 1"
    assert ws.cell(_find(ws, meta, sub="A", after="1.2"), mx.C_STATUS).value in ("Not Submitted", "Overdue")
    groups = [ws.cell(r, 1).value for r in _rows(ws, meta) if str(ws.cell(r, 1).value).startswith("M")]
    assert groups[0].startswith("M1  Design") and "1/3 complete" in groups[0]

    out = wb["Outstanding Items"]
    tags = [out.cell(r, 4).value for r in range(2, out.max_row + 1)]
    assert tags == ["72", "80"]  # 已批准的 71 不在缺件明细里


def test_overdue_shown_for_past_due_not_submitted(session, project, setup):
    data = mx.export_current(session, project["id"], today=date(2026, 10, 24))
    _, ws, meta = _book(data)
    assert ws.cell(_find(ws, meta, sub="A", after="1.2"), mx.C_STATUS).value == "Overdue"
    assert ws.cell(_find(ws, meta, "2.1"), mx.C_STATUS).value == "Not Submitted"  # 没有截止日期
    early = mx.export_current(session, project["id"], today=date(2026, 1, 1))
    _, ws, meta = _book(early)
    assert ws.cell(_find(ws, meta, sub="A", after="1.2"), mx.C_STATUS).value == "Not Submitted"


def test_internal_remarks_never_exported(api_client, project, setup, session):
    d = session.scalars(select(Deliverable).where(
        Deliverable.document_type_id == setup["docs"]["1.1"]["id"])).one()
    d.remarks = "INTERNAL ONLY"
    session.commit()
    for fmt in ("xlsx", "pdf"):
        data = api_client.get(_url(project, f"mdl/export?format={fmt}")).content
        assert b"INTERNAL ONLY" not in data
    wb = load_workbook(io.BytesIO(api_client.get(_url(project, "mdl/export")).content))
    texts = [str(c.value) for w in wb.worksheets for row in w.iter_rows() for c in row if c.value]
    assert not any("INTERNAL ONLY" in t or "Internal checklist" in t for t in texts)


def test_chinese_excel(api_client, project, setup):
    wb, ws, meta = _book(api_client.get(_url(project, "mdl/export?lang=zh")).content)
    assert wb.sheetnames[:3] == ["MDL", "版本记录", "缺件明细"]
    assert ws.cell(_find(ws, meta, "1.1"), mx.C_STATUS).value == "已批准"


# ---------------------------------------------------------------- PDF


@pytest.mark.parametrize("lang", ["en", "zh"])
def test_pdf_export(api_client, project, setup, lang):
    resp = api_client.get(_url(project, f"mdl/export?format=pdf&lang={lang}"))
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF") and len(resp.content) > 3000
    assert rb"/Author (docbase \(system generated\))" in resp.content  # PDF 文件属性里的作者


# ---------------------------------------------------------------- 发出与版本


def test_issue_records_versions_and_marks_changes(api_client, project, setup):
    resp = api_client.post(_url(project, "mdl/issue"), json={"description": "First issue"},
                           headers=USER)
    assert (resp.headers["X-MDL-Revision"], resp.headers["X-MDL-New-Revision"]) == ("Rev.0", "true")
    wb, ws, meta = _book(resp.content)
    assert ws["H3"].value == "Rev.0"
    history = wb["Revision History"]
    assert (history["A2"].value, history["C2"].value) == ("Rev.0", "First issue")

    resp = api_client.post(_url(project, "mdl/issue"), json={"format": "pdf"}, headers=USER)
    assert (resp.headers["X-MDL-Revision"], resp.headers["X-MDL-New-Revision"]) == ("Rev.0", "false")
    assert resp.content.startswith(b"%PDF")

    _call(api_client, "POST", project, "document-types/bulk",
          {"milestone_id": setup["m1"]["id"], "rows": [{"name": "Noise report"}]}, 201)
    _call(api_client, "PATCH", project, f"document-types/{setup['docs']['1.1']['id']}",
          {"name": "Spec & Data (rev)"})
    resp = api_client.post(_url(project, "mdl/issue"), json={"description": "Client letter CL-1"},
                           headers=USER)
    assert resp.headers["X-MDL-Revision"] == "Rev.1"
    wb, ws, meta = _book(resp.content)
    chg = {ws.cell(r, mx.C_NO).value: ws.cell(r, mx.C_CHG).value for r in _rows(ws, meta)
           if ws.cell(r, mx.C_NO).value}
    assert chg["1.5"] == "NEW" and chg["1.1"] == "CHG" and not chg["1.2"]
    assert wb["Revision History"].max_row == 3  # Rev.0、Rev.1

    _, ws, meta = _book(api_client.get(_url(project, "mdl-versions/0/export")).content)
    titles = [ws.cell(r, mx.C_TITLE).value for r in _rows(ws, meta)]
    assert "Noise report" not in titles  # 结构用那一版的


# ---------------------------------------------------------------- 导回


def _edit_and_preview(api_client, project, edit):
    wb, ws, meta = _book(api_client.get(_url(project, "mdl/export")).content, data_only=False)
    edit(ws, meta)
    buffer = io.BytesIO()
    wb.save(buffer)
    resp = api_client.post(_url(project, "mdl/import/preview"),
                           files={"file": ("returned.xlsx", buffer.getvalue())}, headers=USER)
    assert resp.status_code == 200, resp.text
    return resp.json()


def test_round_trip(api_client, project, setup, session):
    def edit(ws, meta):
        r21 = _find(ws, meta, "2.1")
        ws.cell(r21, mx.C_STATUS, "Under Review")
        ws.cell(r21, mx.C_SDATE, date(2026, 10, 20))
        ws.cell(r21, mx.C_SUB, date(2026, 10, 15))
        ws.cell(_find(ws, meta, sub="B", after="1.2"), mx.C_STATUS, "N/A")
        ws.cell(_find(ws, meta, sub="A", after="1.2"), mx.C_REM, "module A note")
        ws.cell(_find(ws, meta, "1.1"), mx.C_REM, "updated client note")
        ws.cell(_find(ws, meta, "1.3"), mx.C_STATUS, "3/3 Approved")  # 汇总行：只读

    preview = _edit_and_preview(api_client, project, edit)
    assert preview["ok"], preview["errors"]
    targets = sorted((c["target"], tuple(sorted(c["changes"]))) for c in preview["changes"])
    assert targets == [
        ("deliverable", ("status",)),
        ("deliverable", ("status", "status_date", "submitted_date")),
        ("document_type", ("mdl_remarks",)),
        ("partition", ("mdl_remarks",)),
    ]
    assert [w["code"] for w in preview["warnings"]] == ["mdl_import.readonly_changed"]

    counts = _call(api_client, "POST", project, "mdl/import/confirm", {"token": preview["token"]})
    assert counts == {"deliverables": 2, "remarks": 2}
    session.expire_all()
    d21 = session.scalars(select(Deliverable).where(
        Deliverable.document_type_id == setup["docs"]["2.1"]["id"])).one()
    assert (d21.status.value, d21.responsible.value) == ("under_review", "client")
    assert (d21.status_date, d21.submitted_date) == (date(2026, 10, 20), date(2026, 10, 15))
    assert d21.status_basis == "MDL 导回（returned.xlsx）"
    events = session.scalars(select(Event).where(Event.project_id == project["id"],
                                                 Event.event_type == "deliverable.status_changed"))
    assert len(list(events)) == 2
    assert _edit_and_preview(api_client, project, lambda ws, meta: None)["changes"] == []


def test_overdue_label_round_trips_as_not_submitted(session, project, setup):
    """导出时显示"逾期"，原样导回不算改动。"""
    data = mx.export_current(session, project["id"], today=date(2026, 10, 24))
    result = mx.parse_return(session, project["id"], data)
    assert result["ok"] and result["changes"] == []


def test_return_rejects_other_files(api_client, project, setup):
    other = api_client.post("/api/projects", json={"code": "MX-2", "name": "other"},
                            headers=USER).json()
    data = api_client.get(_url(other, "mdl/export")).content
    resp = api_client.post(_url(project, "mdl/import/preview"),
                           files={"file": ("x.xlsx", data)}, headers=USER).json()
    assert resp["errors"][0]["code"] == "mdl_import.other_project"
    buffer = io.BytesIO()
    Workbook().save(buffer)
    resp = api_client.post(_url(project, "mdl/import/preview"),
                           files={"file": ("y.xlsx", buffer.getvalue())}, headers=USER).json()
    assert resp["errors"][0]["code"] == "mdl_import.not_platform_file"


def test_bad_status_blocks_confirm(api_client, project, setup):
    preview = _edit_and_preview(api_client, project, lambda ws, meta: ws.cell(
        _find(ws, meta, "2.1"), mx.C_STATUS, "Maybe"))
    assert not preview["ok"] and preview["errors"][0]["code"] == "mdl_import.bad_status"
    resp = api_client.post(_url(project, "mdl/import/confirm"), json={"token": preview["token"]},
                           headers=USER)
    assert resp.json()["detail"]["code"] == "mdl_import.has_errors"


# ---------------------------------------------------------------- 项目设置、修改应交项


def test_project_mdl_header_fields(api_client, project):
    assert (project["client_name"], project["mdl_doc_no"]) == ("Client Co.", "MX-1-MDL-0001")
    updated = _call(api_client, "PATCH", project, "", {"mdl_checked_by": "PE", "mdl_approved_by": "PM"})
    assert (updated["mdl_checked_by"], updated["mdl_approved_by"]) == ("PE", "PM")


def test_update_deliverable_rules(project, setup, session):
    d = session.scalars(select(Deliverable).where(
        Deliverable.document_type_id == setup["docs"]["1.1"]["id"])).one()
    assert d.status.value == "approved"
    d.status_date = date(2026, 3, 1)
    session.flush()
    update_deliverable(session, project["id"], d.id, {"status": "revise_resubmit"},
                       Actor(name="alice"), reason="客户函 CL-9")
    assert (d.responsible.value, d.status_date, d.status_basis) == ("ours", None, "客户函 CL-9")
    update_deliverable(session, project["id"], d.id, {"status": "submitted"}, Actor(name="alice"))
    assert d.responsible.value == "ours"
